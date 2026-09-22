"""Local bilingual product-asset retrieval with country-stock filtering."""

from __future__ import annotations

import argparse
import contextlib
import gc
import hashlib
import importlib.util
import json
import math
import os
import sqlite3
import sys
from collections import defaultdict
from collections.abc import Iterable, Iterator
from pathlib import Path
from time import perf_counter
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen

import numpy as np
from openpyxl import load_workbook

from .normalize import keyword_tokens
from .storage import _fts_text


DEFAULT_ASSET_DB = Path(r"E:\Project\store-assortment-copilot\var\product-asset\catalog.sqlite3")
DEFAULT_STOCK = Path(
    r"E:\download\Chrome下载\真仓库存明细数据-普通商品-汇总数据-1789685524432.xlsx"
)
DEFAULT_QDRANT = "http://127.0.0.1:6333"
DEFAULT_COLLECTION = "product_asset_bilingual_1024_v1"
COUNTRIES = {
    "PH": "PH", "菲律宾": "PH",
    "TH": "TH", "泰国": "TH",
    "VN": "VN", "越南": "VN",
    "MY": "MY", "马来西亚": "MY",
}
RRF_OFFSET = 60


def load_products(path: Path) -> dict[str, dict[str, object]]:
    connection = sqlite3.connect(f"file:{Path(path).as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        rows = connection.execute(
            """SELECT main_sku, standard_name_cn, standard_name_en, inventory_metadata_json
               FROM products
               WHERE active=1 AND indexable=1
               ORDER BY main_sku COLLATE NOCASE"""
        ).fetchall()
    finally:
        connection.close()
    products = {row["main_sku"]: dict(row) for row in rows}
    if not products or any(
        not item["standard_name_cn"].strip() or not item["standard_name_en"].strip()
        for item in products.values()
    ):
        raise ValueError("product asset needs active/indexable main SKUs with both cleaned names")
    return products


def child_main_map(products: dict[str, dict[str, object]]) -> dict[str, str]:
    result: dict[str, str] = {}
    for main_sku, product in products.items():
        try:
            child_skus = json.loads(str(product["inventory_metadata_json"]))["child_skus"]
        except (json.JSONDecodeError, KeyError, TypeError) as error:
            raise ValueError(f"invalid child SKU metadata for {main_sku}") from error
        if not isinstance(child_skus, list) or not all(
            isinstance(child_sku, str) and child_sku.strip() for child_sku in child_skus
        ):
            raise ValueError(f"invalid child SKU metadata for {main_sku}")
        for child_sku in child_skus:
            previous = result.setdefault(child_sku.strip(), main_sku)
            if previous != main_sku:
                raise ValueError(f"child SKU maps to multiple main SKUs: {child_sku}")
    return result


def build_fts(products: dict[str, dict[str, object]]) -> sqlite3.Connection:
    connection = sqlite3.connect(":memory:")
    connection.execute(
        "CREATE VIRTUAL TABLE products_fts USING fts5(main_sku UNINDEXED, cn, en)"
    )
    connection.executemany(
        "INSERT INTO products_fts(main_sku, cn, en) VALUES (?, ?, ?)",
        (
            (sku, _fts_text((item["standard_name_cn"],)), _fts_text((item["standard_name_en"],)))
            for sku, item in products.items()
        ),
    )
    return connection


def keyword_search(
    connection: sqlite3.Connection, language: str, query: str, limit: int
) -> list[tuple[str, float]]:
    if language not in {"cn", "en"}:
        raise ValueError("language must be cn or en")
    tokens = keyword_tokens(query)
    if not tokens:
        return []
    quoted = " OR ".join(f'"{token.replace(chr(34), chr(34) * 2)}"' for token in tokens)
    match = f"{language} : ({quoted})"
    weights = "0.0, 1.0, 0.0" if language == "cn" else "0.0, 0.0, 1.0"
    return [
        (row[0], float(row[1]))
        for row in connection.execute(
            f"""SELECT main_sku, -bm25(products_fts, {weights}) AS score
                FROM products_fts
                WHERE products_fts MATCH ?
                ORDER BY score DESC, main_sku ASC
                LIMIT ?""",
            (match, limit),
        )
    ]


def qdrant_search(
    base_url: str,
    collection: str,
    language: str,
    vector: np.ndarray,
    limit: int,
) -> list[tuple[str, float]]:
    query = np.asarray(vector, dtype=np.float32)
    if query.shape != (1024,) or not np.isfinite(query).all():
        raise ValueError("query vector must contain 1024 finite values")
    payload = json.dumps(
        {
            "query": query.tolist(),
            "using": language,
            "limit": limit,
            "with_payload": ["main_sku"],
        },
        separators=(",", ":"),
    ).encode("utf-8")
    request = Request(
        f"{base_url.rstrip('/')}/collections/{quote(collection, safe='')}/points/query",
        data=payload,
        method="POST",
        headers={"Content-Type": "application/json"},
    )
    try:
        with urlopen(request, timeout=30) as response:
            body = json.loads(response.read().decode("utf-8"))
    except HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Qdrant query failed with HTTP {error.code}: {detail}") from error
    except (OSError, URLError, TimeoutError, json.JSONDecodeError) as error:
        raise RuntimeError(f"Qdrant query failed: {error}") from error
    points = body.get("result", {}).get("points") if isinstance(body, dict) else None
    if body.get("status") != "ok" or not isinstance(points, list):
        raise RuntimeError(f"Qdrant returned an unexpected response: {body!r}")
    hits = []
    for point in points:
        point_payload = point.get("payload") if isinstance(point, dict) else None
        main_sku = point_payload.get("main_sku") if isinstance(point_payload, dict) else None
        if not isinstance(main_sku, str) or not main_sku:
            raise RuntimeError("Qdrant point is missing main_sku payload")
        hits.append((main_sku, float(point["score"])))
    return hits


def fuse_rankings(
    products: dict[str, dict[str, str]],
    rankings: list[tuple],
    limit: int,
) -> list[dict[str, object]]:
    fused: dict[str, dict[str, object]] = {}
    for ranking in rankings:
        channel, query, hits = ranking[:3]
        weight = float(ranking[3]) if len(ranking) == 4 else 1.0
        if weight <= 0:
            raise ValueError("ranking weight must be positive")
        seen: set[str] = set()
        for rank, (main_sku, raw_score) in enumerate(hits, start=1):
            if main_sku in seen or main_sku not in products:
                continue
            seen.add(main_sku)
            item = fused.setdefault(
                main_sku,
                {"rrf_score": 0.0, "channels": set(), "matched_queries": set(), "evidence": []},
            )
            item["rrf_score"] = float(item["rrf_score"]) + weight / (RRF_OFFSET + rank)
            item["channels"].add(channel)
            item["matched_queries"].add(query)
            item["evidence"].append(
                {"channel": channel, "query": query, "query_weight": weight,
                 "rank": rank, "raw_score": raw_score}
            )
    ordered = sorted(fused.items(), key=lambda pair: (-float(pair[1]["rrf_score"]), pair[0]))
    results = []
    for rank, (main_sku, item) in enumerate(ordered[:limit], start=1):
        product = products[main_sku]
        results.append(
            {
                "rank": rank,
                "main_sku": main_sku,
                "standard_name_cn": product["standard_name_cn"],
                "standard_name_en": product["standard_name_en"],
                "rrf_score": round(float(item["rrf_score"]), 8),
                "channels": sorted(item["channels"]),
                "matched_queries": sorted(item["matched_queries"]),
                "evidence": sorted(item["evidence"], key=lambda row: (row["channel"], row["rank"])),
            }
        )
    return results


def _number(value: object, row_number: int, field: str) -> float:
    if value in (None, ""):
        return 0.0
    try:
        result = float(str(value).replace(",", ""))
    except (TypeError, ValueError) as error:
        raise ValueError(f"stock row {row_number}: invalid {field} {value!r}") from error
    if not math.isfinite(result):
        raise ValueError(f"stock row {row_number}: invalid {field} {value!r}")
    return result


def country_code(country: str) -> str | None:
    """The two-letter code, whether the caller wrote "泰国" or "TH"."""
    raw = country.strip()
    return COUNTRIES.get(raw.upper(), COUNTRIES.get(raw))


# How many disagreeing rows the note names before it stops and counts the rest.
SKIPPED_ROWS_SHOWN = 5


def stock_rows(
    path: Path, country: str, allowed_children: dict[str, str],
    disagreeing: list[tuple[int, str, str, str]] | None = None,
) -> Iterator[tuple[str, str, float]]:
    """Each in-country stock row that maps to a known child, as main, child, quantity.

    One pass over the workbook, so the caller decides what to fold it into: the
    recall needs totals per main SKU, and the exported sub-SKU sheet needs the
    same quantities kept by the child the goods are actually in. Rows whose child
    the catalogue does not know, or that belong to another country, are dropped
    rather than guessed at.

    A row whose 主SKU column contradicts the catalogue's own mapping is dropped
    too, and handed to ``disagreeing`` so the caller can say so. Refusing to run
    would mean one bad line in a 200,000-line export costs a whole pass, and the
    operator has no way to fix a file they cannot see the fault in.
    """
    code = country_code(country)
    if code is None:
        raise ValueError(f"unsupported country: {country!r}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["汇总表格"]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(next(rows))
        required = ("子SKU", "主SKU", "国家", "库存中心库存", "公共池库存")
        if any(headers.count(header) != 1 for header in required):
            raise ValueError("stock needs unique headers: " + ", ".join(required))
        positions = {header: headers.index(header) for header in required}
        for row_number, row in enumerate(rows, start=2):
            child_sku = str(row[positions["子SKU"]] or "").strip()
            main_sku = str(row[positions["主SKU"]] or "").strip()
            expected_main_sku = allowed_children.get(child_sku)
            if expected_main_sku is None:
                continue
            if main_sku != expected_main_sku:
                if disagreeing is not None:
                    disagreeing.append((row_number, child_sku, main_sku, expected_main_sku))
                continue
            raw_country = str(row[positions["国家"]] or "").strip()
            if COUNTRIES.get(raw_country.upper(), COUNTRIES.get(raw_country)) != code:
                continue
            yield main_sku, child_sku, _number(
                row[positions["库存中心库存"]], row_number, "库存中心库存"
            ) + _number(row[positions["公共池库存"]], row_number, "公共池库存")
    finally:
        workbook.close()


def read_country_stock(
    path: Path, country: str, allowed_children: dict[str, str],
    notes: list[str] | None = None,
) -> tuple[dict[str, float], dict[str, float]]:
    """Country stock by main SKU and by child SKU, out of the one read.

    The second map keeps only what this country can actually ship, because a
    child with no row or a row summing to nothing means the same thing to
    whoever has to pick one.
    """
    by_main: defaultdict[str, float] = defaultdict(float)
    by_child: defaultdict[str, float] = defaultdict(float)
    disagreeing: list[tuple[int, str, str, str]] = []
    for main_sku, child_sku, quantity in stock_rows(path, country, allowed_children, disagreeing):
        by_main[main_sku] += quantity
        if quantity > 0:
            by_child[child_sku] += quantity
    if notes is not None:
        notes.extend(_disagreement_note(disagreeing))
    return dict(by_main), dict(by_child)


def _disagreement_note(disagreeing: list[tuple[int, str, str, str]]) -> list[str]:
    """Name the rows the two files disagree about, without pasting in hundreds.

    A systematic mapping error would otherwise write one line per row into a
    field the operator reads on the page, and bury the number that says how
    widespread the damage is.
    """
    if not disagreeing:
        return []
    lines = [f"库存文件里有 {len(disagreeing)} 行与产品库的子款归属对不上，这些行已跳过："]
    lines += [f"第 {number} 行 {child}：产品库里挂在 {expected} 下，文件里记的是 {claimed}。"
              for number, child, claimed, expected in disagreeing[:SKIPPED_ROWS_SHOWN]]
    if len(disagreeing) > SKIPPED_ROWS_SHOWN:
        lines.append(f"其余 {len(disagreeing) - SKIPPED_ROWS_SHOWN} 行同样跳过。")
    return lines


def load_country_stock(path: Path, country: str, allowed_children: dict[str, str]) -> dict[str, float]:
    return read_country_stock(path, country, allowed_children)[0]


def load_inventory_context(
    path: Path, country: str, allowed_children: dict[str, str]
) -> tuple[str, str, dict[str, float] | None]:
    raw = country.strip()
    code = country_code(raw)
    if code is None:
        return raw.upper(), "unavailable", None
    return code, "available", load_country_stock(path, code, allowed_children)


def load_children(asset_db: Path, main_skus: Iterable[str]) -> dict[str, list[tuple[str, str]]]:
    """Each main SKU's children, with the name the catalogue holds for them.

    Read from the catalogue and not from the stock workbook: the workbook only
    knows the children that appear in it, and the sheet has to list every child
    the product has, including the ones this country cannot ship.
    """
    wanted = sorted({str(sku) for sku in main_skus})
    children: dict[str, list[tuple[str, str]]] = {sku: [] for sku in wanted}
    if not wanted:
        return children
    connection = sqlite3.connect(f"file:{Path(asset_db).as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        for start in range(0, len(wanted), 500):
            batch = wanted[start:start + 500]
            rows = connection.execute(
                f"SELECT main_sku, inventory_metadata_json, raw_evidence_json FROM products "
                f"WHERE active=1 AND indexable=1 AND main_sku IN ({','.join('?' * len(batch))})",
                batch,
            ).fetchall()
            for row in rows:
                names = _child_names(row["raw_evidence_json"])
                for child in _metadata_children(row["inventory_metadata_json"]):
                    children[row["main_sku"]].append((child, names.get(child, "")))
    finally:
        connection.close()
    return children


def _metadata_children(value: object) -> list[str]:
    try:
        child_skus = json.loads(str(value))["child_skus"]
    except (json.JSONDecodeError, KeyError, TypeError):
        return []
    if not isinstance(child_skus, list):
        return []
    return [str(child).strip() for child in child_skus if str(child).strip()]


def _child_names(value: object) -> dict[str, str]:
    """The child's own product name, off the row the catalogue was built from."""
    try:
        records = json.loads(str(value))
    except (json.JSONDecodeError, TypeError):
        return {}
    if not isinstance(records, list):
        return {}
    names: dict[str, str] = {}
    for record in records:
        if not isinstance(record, dict):
            continue
        raw = record.get("raw")
        child = str(record.get("child_sku") or "").strip()
        name = str(raw.get("商品名称") or "").strip() if isinstance(raw, dict) else ""
        if child and name and child not in names:
            names[child] = name
    return names


def filter_country(
    global_candidates: list[dict[str, object]], stock: dict[str, float] | None, limit: int
) -> list[dict[str, object]]:
    if stock is None:
        for candidate in global_candidates:
            candidate["country_available_quantity"] = None
            candidate["country_available"] = None
        return []
    results = []
    for candidate in global_candidates:
        quantity = stock.get(str(candidate["main_sku"]), 0.0)
        candidate["country_available_quantity"] = round(quantity, 4) if quantity > 0 else None
        candidate["country_available"] = quantity > 0
        if quantity > 0 and len(results) < limit:
            results.append(dict(candidate))
    return results


def embed_queries(asset_db: Path, queries: dict[str, list[str]]) -> dict[tuple[str, str], np.ndarray]:
    script = asset_db.parent / "vectorize_bilingual_1024.py"
    spec = importlib.util.spec_from_file_location("product_asset_vectorizer", script)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load vectorizer: {script}")
    module = importlib.util.module_from_spec(spec)
    with contextlib.redirect_stdout(sys.stderr):
        spec.loader.exec_module(module)
    os.environ.setdefault("HF_HUB_OFFLINE", "1")
    os.environ.setdefault("TRANSFORMERS_OFFLINE", "1")
    vectors: dict[tuple[str, str], np.ndarray] = {}
    for language in ("cn", "en"):
        texts = queries.get(language, [])
        if not texts:
            continue
        with contextlib.redirect_stdout(sys.stderr):
            tokenizer, model, device, dtype, _ = module.load_model(module.MODELS[language])
        batch = module.encode_batch(tokenizer, model, device, dtype, texts)
        vectors.update(((language, text), vector) for text, vector in zip(texts, batch, strict=True))
        del model, tokenizer, batch
        gc.collect()
        if module.torch.cuda.is_available():
            module.torch.cuda.empty_cache()
    return vectors


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def run(args: argparse.Namespace) -> dict[str, object]:
    started = perf_counter()
    queries = {
        "cn": list(dict.fromkeys(text.strip() for text in args.cn if text.strip())),
        "en": list(dict.fromkeys(text.strip() for text in args.en if text.strip())),
    }
    if not queries["cn"] and not queries["en"]:
        raise ValueError("provide at least one --cn or --en query")
    products = load_products(args.asset_db)
    fts = build_fts(products)
    try:
        vectors = embed_queries(args.asset_db, queries)
        rankings: list[tuple[str, str, list[tuple[str, float]]]] = []
        for language in ("cn", "en"):
            for query in queries[language]:
                rankings.append(
                    (f"{language}_keyword", query, keyword_search(fts, language, query, args.channel_limit))
                )
                rankings.append(
                    (
                        f"{language}_vector",
                        query,
                        qdrant_search(
                            args.qdrant,
                            args.collection,
                            language,
                            vectors[(language, query)],
                            args.channel_limit,
                        ),
                    )
                )
    finally:
        fts.close()
    global_candidates = fuse_rankings(products, rankings, args.global_limit)
    country, inventory_coverage, stock = load_inventory_context(
        args.stock, args.country, child_main_map(products)
    )
    country_candidates = filter_country(global_candidates, stock, args.country_limit)
    return {
        "result_stage": "retrieval_pilot_not_reranked",
        "queries": queries,
        "country": country,
        "inventory_coverage": inventory_coverage,
        "retrieval": {
            "fusion": f"RRF(k={RRF_OFFSET})",
            "channel_limit": args.channel_limit,
            "global_limit": args.global_limit,
            "country_limit": args.country_limit,
            "global_candidates_are_preserved": True,
            "country_filter": "sum(库存中心库存 + 公共池库存) > 0"
            if inventory_coverage == "available" else None,
        },
        "assets": {
            "product_db": str(args.asset_db.resolve()),
            "product_count": len(products),
            "stock_file": str(args.stock.resolve()),
            "stock_sha256": _sha256(args.stock) if inventory_coverage == "available" else None,
            "qdrant": args.qdrant,
            "collection": args.collection,
        },
        "global_candidates": global_candidates,
        "country_candidates": country_candidates,
        "elapsed_seconds": round(perf_counter() - started, 3),
    }


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__)
    result.add_argument("--cn", action="append", default=[], help="Chinese product need; repeatable")
    result.add_argument("--en", action="append", default=[], help="English product need; repeatable")
    result.add_argument("--country", required=True, help="PH/TH/VN/MY or Chinese country name")
    result.add_argument("--asset-db", type=Path, default=DEFAULT_ASSET_DB)
    result.add_argument("--stock", type=Path, default=DEFAULT_STOCK)
    result.add_argument("--qdrant", default=DEFAULT_QDRANT)
    result.add_argument("--collection", default=DEFAULT_COLLECTION)
    result.add_argument("--channel-limit", type=int, default=200)
    result.add_argument("--global-limit", type=int, default=200)
    result.add_argument("--country-limit", type=int, default=40)
    result.add_argument("--output", type=Path)
    return result


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    if min(args.channel_limit, args.global_limit, args.country_limit) <= 0:
        raise ValueError("retrieval limits must be positive")
    result = run(args)
    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(output + "\n", encoding="utf-8")
        print(args.output.resolve())
    else:
        print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
