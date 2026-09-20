"""Read a complete country-stock snapshot; never sum salesperson duplicates."""

from collections import defaultdict
from datetime import UTC, datetime
from itertools import zip_longest
from pathlib import Path
import math

from openpyxl import load_workbook

from .eligibility import is_excluded_product
from .versioning import _sha256_file

COUNTRIES = {"菲律宾": "PH", "泰国": "TH", "越南": "VN", "马来西亚": "MY"}


def read_stock_snapshot(path: Path, sku_to_main: dict[str, str], *, allow_unmatched: bool = False) -> dict:
    source = Path(path)
    fingerprint = _sha256_file(source)
    workbook = load_workbook(source, read_only=True, data_only=True)
    available = defaultdict(set)
    available_children = set()
    detail_count = 0
    excluded_positive_count = 0
    known_mains = set(sku_to_main.values())
    main_only_children = set()
    unmatched = []
    try:
        sheet = workbook["汇总表格"]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        required = ("子SKU", "主SKU", "国家", "海外仓可发")
        if any(headers.count(key) != 1 for key in required):
            raise ValueError("stock needs unique headers: " + ", ".join(required))
        positions = {key: headers.index(key) for key in required}
        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            def value(key):
                index = positions[key]
                return row[index] if index < len(row) else None

            sku = str(value("子SKU") or "").strip()
            if not sku and str(row[0] or "").strip() == "汇总":
                continue
            if not sku:
                raise ValueError(f"stock row {row_number}: missing child SKU")
            country_raw = str(value("国家") or "").strip()
            country = COUNTRIES.get(country_raw, country_raw)
            if country not in COUNTRIES.values():
                raise ValueError(f"stock row {row_number}: unknown country {country_raw!r}")
            available[country]  # Keep countries that currently have no available stock.
            detail_count += 1
            quantity = float(value("海外仓可发") or 0)
            if not math.isfinite(quantity):
                raise ValueError(f"stock row {row_number}: invalid available quantity")
            if quantity <= 0:
                continue
            reported_main = str(value("主SKU") or "").strip()
            mapped_main = sku_to_main.get(sku)
            if mapped_main and reported_main and reported_main != mapped_main:
                raise ValueError(f"stock row {row_number}: main SKU disagrees with ERP for {sku!r}")
            main_sku = reported_main or mapped_main
            if is_excluded_product(main_sku, sku):
                excluded_positive_count += 1
                continue
            if not main_sku or main_sku not in known_mains:
                if allow_unmatched and main_sku:
                    # Preserve source evidence, but do not invent an ERP identity
                    # or silently turn unmatched stock into zero inventory.
                    unmatched.append({
                        "source_row": row_number, "country": country,
                        "main_sku": main_sku, "child_sku": sku,
                        "raw": {header: (cell.isoformat() if isinstance(cell, datetime) else cell)
                                for header, cell in zip_longest(headers, row) if isinstance(header, str) and header},
                        "reason": "main_sku_missing_from_erp",
                    })
                    continue
                raise ValueError(f"stock row {row_number}: main SKU {main_sku!r} missing from ERP (child {sku!r})")
            if mapped_main is None:
                # Inventory already supplies the exact main-SKU identity. A new
                # child variant must not block a known family; no fuzzy matching.
                main_only_children.add(sku)
            available[country].add(main_sku)
            available_children.add(sku)
    finally:
        workbook.close()
    if detail_count == 0:
        raise ValueError("stock snapshot has no detail rows")
    if _sha256_file(source) != fingerprint:
        raise ValueError("stock file changed during import")
    return {
        "source": str(source.resolve()),
        "source_sha256": fingerprint,
        "source_modified_at": datetime.fromtimestamp(source.stat().st_mtime, UTC).isoformat(),
        "imported_at": datetime.now(UTC).isoformat(),
        "stock_basis": "海外仓可发 > 0; same-country sharing; no quantity summation",
        "detail_rows": detail_count,
        "excluded_positive_rows": excluded_positive_count,
        "available_child_count": len(available_children),
        "main_only_matched_child_count": len(main_only_children),
        "erp_mapping_complete": not unmatched,
        "unmatched_positive_rows": unmatched,
        "unmatched_positive_main_skus": sorted({row["main_sku"] for row in unmatched}),
        "country_main_skus": {country: sorted(mains) for country, mains in sorted(available.items())},
    }
