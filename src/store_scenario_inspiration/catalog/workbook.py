"""Read ERP catalog workbooks without modifying their source files."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import SourceRow


REQUIRED_HEADERS = frozenset(
    {
        "sku",
        "主SKU",
        "商品名称",
        "英文名称",
        "英文关键字",
        "销售状态",
        "商品目录",
        "商品一级目录",
        "商品二级目录",
        "商品三级目录",
        "商品四级目录",
    }
)


class CatalogSchemaError(ValueError):
    """Raised when a workbook sheet does not match the ERP row schema."""


def _reset_dimensions(worksheet: Worksheet) -> None:
    reset = getattr(worksheet, "reset_dimensions", None)
    if reset is not None:
        reset()


def _header_map(worksheet: Worksheet) -> dict[str, int]:
    first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {
        value: index
        for index, value in enumerate(first_row)
        if isinstance(value, str) and value
    }


def _missing_headers(header_map: dict[str, int]) -> tuple[str, ...]:
    return tuple(sorted(REQUIRED_HEADERS.difference(header_map)))


def _select_sheets(workbook: object, sheet_name: str | None) -> tuple[Worksheet, ...]:
    if sheet_name is not None:
        try:
            worksheet = workbook[sheet_name]  # type: ignore[index]
        except KeyError as error:
            raise CatalogSchemaError(f"worksheet {sheet_name!r} does not exist") from error
        _reset_dimensions(worksheet)
        missing = _missing_headers(_header_map(worksheet))
        if missing:
            raise CatalogSchemaError(
                f"worksheet {sheet_name!r} is missing required headers: {', '.join(missing)}"
            )
        return (worksheet,)

    matches = []
    for worksheet in workbook.worksheets:  # type: ignore[attr-defined]
        _reset_dimensions(worksheet)
        if not _missing_headers(_header_map(worksheet)):
            matches.append(worksheet)
    if not matches:
        available_names = ", ".join(
            worksheet.title for worksheet in workbook.worksheets  # type: ignore[attr-defined]
        )
        raise CatalogSchemaError(
            "expected at least one worksheet containing all required headers; "
            f"found none (available: {available_names or 'none'})"
        )
    return tuple(matches)


def _source_text(value: object) -> str:
    return "" if value is None else str(value)


def iter_source_rows(path: Path, sheet_name: str | None = None) -> Iterator[SourceRow]:
    """Yield raw string rows from schema-matching workbook sheets.

    The workbook is opened with openpyxl's read-only, cached-value mode and is
    always closed after iteration.  This function never saves or otherwise
    mutates the supplied source workbook.
    """

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in _select_sheets(workbook, sheet_name):
            headers = _header_map(worksheet)
            for values in worksheet.iter_rows(min_row=2, values_only=True):
                if not any(value is not None for value in values):
                    continue

                def value_for(header: str) -> str:
                    index = headers[header]
                    value = values[index] if index < len(values) else None
                    return _source_text(value)

                yield SourceRow(
                    sku=value_for("sku"),
                    main_sku=value_for("主SKU"),
                    product_name=value_for("商品名称"),
                    english_name=value_for("英文名称"),
                    english_keywords=value_for("英文关键字"),
                    sales_status_raw=value_for("销售状态"),
                    product_catalog=value_for("商品目录"),
                    category_level_1=value_for("商品一级目录"),
                    category_level_2=value_for("商品二级目录"),
                    category_level_3=value_for("商品三级目录"),
                    category_level_4=value_for("商品四级目录"),
                )
    finally:
        workbook.close()
