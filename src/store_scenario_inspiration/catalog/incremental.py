"""Durable daily catalog imports on top of an existing SQLite catalog."""

from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import tempfile
from dataclasses import asdict, replace
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .build import build_documents
from .eligibility import is_excluded_main_sku, is_excluded_product
from .models import BuildManifest, ChildVariant, ProductFamilyDocument, SourceRow
from .normalize import clean_optional_text, normalize_compare, make_vector_text
from .storage import _document_from_json, _fts_text, _json, _manifest_from_json
from .workbook import REQUIRED_HEADERS, _header_map, _select_sheets


_DAILY_SCHEMA = """
CREATE TABLE IF NOT EXISTS daily_meta (
    id INTEGER PRIMARY KEY CHECK (id = 1),
    baseline_version_id TEXT NOT NULL,
    baseline_source_sha256 TEXT NOT NULL,
    revision INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS daily_source_rows (
    child_sku TEXT PRIMARY KEY,
    main_sku TEXT NOT NULL,
    source_sheet TEXT NOT NULL,
    source_row INTEGER NOT NULL,
    raw_json TEXT NOT NULL,
    row_hash TEXT NOT NULL,
    evidence_json TEXT NOT NULL,
    evidence_hash TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS daily_imports (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    source_path TEXT NOT NULL,
    source_hash TEXT NOT NULL,
    applied INTEGER NOT NULL,
    report_json TEXT NOT NULL,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS daily_name_review (
    main_sku TEXT PRIMARY KEY,
    evidence_hash TEXT NOT NULL,
    evidence_json TEXT NOT NULL,
    status TEXT NOT NULL,
    cn_name TEXT NOT NULL DEFAULT '',
    en_name TEXT NOT NULL DEFAULT '',
    aliases_json TEXT NOT NULL DEFAULT '[]',
    reason TEXT NOT NULL DEFAULT ''
);
"""
_SEMANTIC_HEADERS = (
    "商品名称",
    "英文名称",
    "英文关键字",
    "商品目录",
    "商品一级目录",
    "商品二级目录",
    "商品三级目录",
    "商品四级目录",
)


def initialize_daily_catalog(baseline: Path, target: Path) -> None:
    manifest = _baseline_manifest(baseline)
    if target.exists():
        connection = _connect_ro(target)
        try:
            if not _table_exists(connection, "daily_meta"):
                raise ValueError("existing daily catalog is missing daily metadata")
            row = connection.execute(
                "SELECT baseline_version_id, baseline_source_sha256 FROM daily_meta WHERE id = 1"
            ).fetchone()
            if row != (manifest.version_id, manifest.source_sha256):
                raise ValueError("existing daily catalog belongs to a different baseline")
        finally:
            connection.close()
        return
    if not target.parent.exists():
        raise FileNotFoundError(f"catalog directory does not exist: {target.parent}")

    temporary_name: str | None = None
    source: sqlite3.Connection | None = None
    target_connection: sqlite3.Connection | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{target.name}.", suffix=".tmp", dir=target.parent, delete=False
        ) as temporary:
            temporary_name = temporary.name
        source = _connect_ro(baseline)
        target_connection = sqlite3.connect(temporary_name)
        source.backup(target_connection)
        source.close()
        source = None
        with target_connection:
            _ensure_daily(target_connection)
            target_connection.execute(
                """INSERT INTO daily_meta(
                    id, baseline_version_id, baseline_source_sha256, revision
                ) VALUES (1, ?, ?, 0)""",
                (manifest.version_id, manifest.source_sha256),
            )
        target_connection.close()
        target_connection = None
        os.link(temporary_name, target)
        Path(temporary_name).unlink()
        temporary_name = None
    except Exception:
        if source is not None:
            source.close()
        if target_connection is not None:
            target_connection.close()
        if temporary_name is not None:
            Path(temporary_name).unlink(missing_ok=True)
        raise


def catalog_revision(path: Path) -> int:
    connection = _connect_ro(path)
    try:
        if not _table_exists(connection, "daily_meta"):
            return 0
        row = connection.execute("SELECT revision FROM daily_meta WHERE id = 1").fetchone()
        return 0 if row is None else int(row[0])
    finally:
        connection.close()


def import_products(path: Path, source: Path, *, apply: bool = False) -> dict[str, Any]:
    source_hash = _sha256(source)
    rows = _read_source(source)
    rows_by_child: dict[str, _ImportRow] = {}
    duplicates: dict[str, str] = {}
    for row in rows:
        if not row.child_sku or not row.main_sku:
            raise ValueError("each source row requires non-empty sku and main_sku")
        old = rows_by_child.get(row.child_sku)
        if old is None:
            rows_by_child[row.child_sku] = row
        elif old.row_hash != row.row_hash:
            duplicates[row.child_sku] = row.child_sku
    if duplicates:
        raise ValueError(f"conflicting duplicate child rows: {', '.join(sorted(duplicates))}")
    rows = tuple(rows_by_child.values())

    connection = _connect_rw(path)
    try:
        _ensure_daily(connection)
        groups = _group_by_main(rows)
        current_docs = _documents(connection, tuple(groups))
        child_to_main = _child_main_map(connection, tuple(row.child_sku for row in rows))
        for row in rows:
            old_main = child_to_main.get(row.child_sku)
            if old_main is not None and old_main != row.main_sku:
                raise ValueError(f"child SKU reparent is not allowed: {row.child_sku}")

        old_source = _source_rows(connection)
        changed_raw = [row for row in rows if old_source.get(row.child_sku, {}).get("row_hash") != row.row_hash]
        merged_docs = dict(current_docs)
        changed_docs: list[ProductFamilyDocument] = []
        for main_sku, incoming in groups.items():
            document = _merge_document(current_docs.get(main_sku), incoming)
            merged_docs[main_sku] = document
            if _json(document.to_index_dict()) != _json(current_docs.get(main_sku).to_index_dict() if current_docs.get(main_sku) else {}):
                changed_docs.append(document)

        pending_reviews = _refresh_reviews(connection, current_docs, merged_docs, rows, old_source)
        revision = catalog_revision(path)
        effective_change = bool(changed_docs)
        report = {
            "source_rows": len(rows),
            "existing_children": sum(1 for row in rows if row.child_sku in child_to_main),
            "new_children": sum(1 for row in rows if row.child_sku not in child_to_main),
            "existing_mains": sum(1 for main in groups if main in current_docs),
            "new_mains": sum(1 for main in groups if main not in current_docs),
            "changed_metadata_count": len(changed_docs),
            "pending_name_mains": pending_reviews,
            "excluded_count": sum(1 for row in rows if is_excluded_product(row.main_sku, row.child_sku)),
            "retained_children": sum(
                len(document.children)
                for main, document in current_docs.items()
                if main in groups
            )
            - sum(1 for row in rows if row.child_sku in child_to_main),
            "retained_mains": _document_count(connection) - sum(1 for main in groups if main in current_docs),
            "revision": revision + (1 if effective_change and apply else 0),
            "applied": apply,
            "file_hash": source_hash,
            "no_delete_policy": True,
        }

        connection.execute("BEGIN")
        try:
            if _sha256(source) != source_hash:
                raise ValueError("source workbook changed during import")
            if apply:
                for row in changed_raw:
                    connection.execute(
                        """INSERT INTO daily_source_rows(
                            child_sku, main_sku, source_sheet, source_row, raw_json,
                            row_hash, evidence_json, evidence_hash
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        ON CONFLICT(child_sku) DO UPDATE SET
                            main_sku = excluded.main_sku,
                            source_sheet = excluded.source_sheet,
                            source_row = excluded.source_row,
                            raw_json = excluded.raw_json,
                            row_hash = excluded.row_hash,
                            evidence_json = excluded.evidence_json,
                            evidence_hash = excluded.evidence_hash""",
                        (
                            row.child_sku,
                            row.main_sku,
                            row.sheet,
                            row.row_number,
                            _json(row.raw),
                            row.row_hash,
                            _json(row.evidence),
                            row.evidence_hash,
                        ),
                    )
                for document in changed_docs:
                    _write_document(connection, document)
                _write_review_rows(connection, current_docs, merged_docs, rows, old_source)
                if effective_change:
                    connection.execute(
                        "UPDATE daily_meta SET revision = revision + 1 WHERE id = 1"
                    )
                connection.execute(
                    "INSERT INTO daily_imports(source_path, source_hash, applied, report_json) VALUES (?, ?, ?, ?)",
                    (str(source), source_hash, 1, _json(report)),
                )
                connection.commit()
            else:
                connection.rollback()
        except Exception:
            connection.rollback()
            raise
        return report
    finally:
        connection.close()


def pending_name_reviews(path: Path) -> list[dict[str, Any]]:
    connection = _connect_ro(path)
    try:
        rows = connection.execute(
            """SELECT main_sku, evidence_hash, evidence_json
               FROM daily_name_review
               WHERE status = 'pending'
               ORDER BY main_sku"""
        ).fetchall()
        docs = _documents(connection, tuple(row[0] for row in rows))
        return [
            {
                "main_sku": main_sku,
                "evidence_hash": evidence_hash,
                "previous_cn_names": list(docs[main_sku].cn_names),
                "previous_en_aliases": list(docs[main_sku].en_aliases),
                "source_evidence": json.loads(evidence_json),
                "children": [asdict(child) for child in docs[main_sku].children],
            }
            for main_sku, evidence_hash, evidence_json in rows
            if main_sku in docs
        ]
    finally:
        connection.close()


def apply_name_reviews(path: Path, reviews: list[dict[str, Any]]) -> dict[str, Any]:
    if not isinstance(reviews, list) or any(not isinstance(review, dict) for review in reviews):
        raise ValueError("reviews must be a JSON list of objects")
    seen = [normalize_compare(review.get("main_sku", "")) for review in reviews]
    if len(seen) != len(set(seen)):
        raise ValueError("duplicate main SKU in reviews")
    connection = _connect_rw(path)
    try:
        _ensure_daily(connection)
        docs = _documents(connection, tuple(seen))
        replacements: list[ProductFamilyDocument] = []
        connection.execute("BEGIN")
        try:
            for review, main_sku in zip(reviews, seen, strict=True):
                if main_sku not in docs:
                    raise ValueError(f"unknown main SKU: {main_sku}")
                if is_excluded_main_sku(main_sku):
                    raise ValueError(f"excluded main SKU cannot be reviewed: {main_sku}")
                stored = connection.execute(
                    """SELECT evidence_hash, status, cn_name, en_name, aliases_json
                       FROM daily_name_review WHERE main_sku = ?""",
                    (main_sku,),
                ).fetchone()
                cn_name = normalize_compare(review.get("cn_name", ""))
                en_name = normalize_compare(review.get("en_name", ""))
                if not isinstance(review.get("aliases", []), list) or any(
                    not isinstance(alias, str) for alias in review.get("aliases", [])
                ):
                    raise ValueError("aliases must be a list of English strings")
                aliases = tuple(
                    normalize_compare(alias) for alias in review.get("aliases", ()) if normalize_compare(alias)
                )
                if not cn_name or not _is_english_text(en_name) or any(
                    not _is_english_text(alias) for alias in aliases
                ):
                    raise ValueError("reviewed Chinese name and English name are required")
                en_aliases = tuple(dict.fromkeys((en_name, *aliases)))
                if stored is None or stored[0] != review.get("evidence_hash"):
                    raise ValueError(f"stale name review evidence for {main_sku}")
                stored_aliases = tuple(json.loads(stored[4]))
                if stored[1] == "reviewed":
                    if (stored[2], stored[3], stored_aliases) == (
                        cn_name,
                        en_name,
                        en_aliases,
                    ):
                        continue
                    raise ValueError(f"stale name review evidence for {main_sku}")
                document = replace(
                    docs[main_sku], cn_names=(cn_name,), en_aliases=en_aliases,
                    vector_text_v1=make_vector_text((cn_name,)),
                    searchable=any(not is_excluded_product(main_sku, child.sku)
                                   for child in docs[main_sku].children),
                    quality_flags=tuple(flag for flag in docs[main_sku].quality_flags
                                        if flag not in {"unsearchable", "missing_english"}),
                )
                replacements.append(document)
                connection.execute(
                    """UPDATE daily_name_review
                       SET status = 'reviewed', cn_name = ?, en_name = ?,
                           aliases_json = ?, reason = ?
                       WHERE main_sku = ?""",
                    (
                        cn_name,
                        en_name,
                        _json(list(en_aliases)),
                        str(review.get("reason", "")),
                        main_sku,
                    ),
                )
            changed = [
                document
                for document in replacements
                if _json(document.to_index_dict()) != _json(docs[document.main_sku].to_index_dict())
            ]
            for document in changed:
                _write_document(connection, document)
            if changed:
                connection.execute("UPDATE daily_meta SET revision = revision + 1 WHERE id = 1")
            connection.commit()
            return {
                "reviewed": len(reviews),
                "changed": len(changed),
                "revision": catalog_revision(path),
            }
        except Exception:
            connection.rollback()
            raise
    finally:
        connection.close()


class _ImportRow:
    def __init__(self, row: SourceRow, raw: dict[str, str], sheet: str, row_number: int) -> None:
        self.row = row
        self.raw = raw
        self.sheet = sheet
        self.row_number = row_number
        self.child_sku = normalize_compare(row.sku)
        self.main_sku = normalize_compare(row.main_sku)
        self.evidence = {
            key: raw[key] for key in _SEMANTIC_HEADERS if clean_optional_text(raw.get(key))
        } if not is_excluded_product(self.main_sku, self.child_sku) else {}
        self.evidence_hash = _hash_json(self.evidence)
        self.row_hash = _hash_json({"raw": raw, "sheet": sheet, "row": row_number})


def _connect_ro(path: Path) -> sqlite3.Connection:
    return sqlite3.connect(Path(path).resolve().as_uri() + "?mode=ro", uri=True)


def _connect_rw(path: Path) -> sqlite3.Connection:
    connection = sqlite3.connect(path)
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def _ensure_daily(connection: sqlite3.Connection) -> None:
    connection.executescript(_DAILY_SCHEMA)


def _table_exists(connection: sqlite3.Connection, name: str) -> bool:
    return connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?", (name,)
    ).fetchone() is not None


def _baseline_manifest(path: Path) -> BuildManifest:
    connection = _connect_ro(path)
    try:
        row = connection.execute(
            "SELECT value_json FROM build_meta WHERE meta_key = 'manifest'"
        ).fetchone()
        if row is None:
            raise ValueError("catalog is missing its build manifest")
        return _manifest_from_json(row[0])
    finally:
        connection.close()


def _documents(
    connection: sqlite3.Connection, main_skus: tuple[str, ...]
) -> dict[str, ProductFamilyDocument]:
    result = {}
    for start in range(0, len(main_skus), 900):
        batch = main_skus[start:start + 900]
        placeholders = ",".join("?" for _ in batch)
        rows = connection.execute(
            f"SELECT main_sku, document_json FROM documents WHERE main_sku IN ({placeholders})",
            batch,
        )
        result.update((row[0], _document_from_json(row[1])) for row in rows)
    return result


def _document_count(connection: sqlite3.Connection) -> int:
    return int(connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0])


def _source_rows(connection: sqlite3.Connection) -> dict[str, dict[str, Any]]:
    if not _table_exists(connection, "daily_source_rows"):
        return {}
    return {
        row[0]: {
            "main_sku": row[1],
            "row_hash": row[2],
            "evidence_json": row[3],
            "evidence_hash": row[4],
        }
        for row in connection.execute(
            "SELECT child_sku, main_sku, row_hash, evidence_json, evidence_hash FROM daily_source_rows"
        )
    }


def _child_main_map(connection: sqlite3.Connection, child_skus: tuple[str, ...]) -> dict[str, str]:
    wanted = set(child_skus)
    result = {}
    for main_sku, sku in connection.execute("SELECT main_sku, sku FROM children"):
        child_sku = normalize_compare(sku)
        if child_sku in wanted:
            if child_sku in result and result[child_sku] != main_sku:
                raise ValueError(f"ambiguous existing child SKU mapping: {child_sku}")
            result[child_sku] = main_sku
    return result


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _hash_json(value: object) -> str:
    return hashlib.sha256(_json(value).encode("utf-8")).hexdigest()


def _is_english_text(value: str) -> bool:
    return bool(value and any("A" <= character.upper() <= "Z" for character in value)
                and not any("\u3400" <= character <= "\u9fff" for character in value))


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def _read_source(path: Path) -> tuple[_ImportRow, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        imported: list[_ImportRow] = []
        for worksheet in _select_sheets(workbook, None):
            headers = _header_map(worksheet)
            header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            header_positions = tuple(
                (index, value)
                for index, value in enumerate(header_row)
                if isinstance(value, str) and value
            )
            seen_headers: set[str] = set()
            repeated_required = []
            for _, header in header_positions:
                if header in REQUIRED_HEADERS and header in seen_headers:
                    repeated_required.append(header)
                seen_headers.add(header)
            if repeated_required:
                raise ValueError(
                    f"worksheet {worksheet.title!r} has duplicate required headers: "
                    + ", ".join(sorted(set(repeated_required)))
                )
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                if not any(value is not None for value in values):
                    continue
                raw = {
                    header: _text(values[index] if index < len(values) else None)
                    for index, header in header_positions
                }

                def value(header: str) -> str:
                    index = headers[header]
                    return _text(values[index] if index < len(values) else None)

                imported.append(
                    _ImportRow(
                        SourceRow(
                            sku=value("sku"),
                            main_sku=value("主SKU"),
                            product_name=value("商品名称"),
                            english_name=value("英文名称"),
                            english_keywords=value("英文关键字"),
                            sales_status_raw=value("销售状态"),
                            product_catalog=value("商品目录"),
                            category_level_1=value("商品一级目录"),
                            category_level_2=value("商品二级目录"),
                            category_level_3=value("商品三级目录"),
                            category_level_4=value("商品四级目录"),
                        ),
                        raw,
                        worksheet.title,
                        row_number,
                    )
                )
        return tuple(imported)
    finally:
        workbook.close()


def _group_by_main(rows: tuple[_ImportRow, ...]) -> dict[str, tuple[_ImportRow, ...]]:
    grouped: dict[str, list[_ImportRow]] = {}
    for row in rows:
        grouped.setdefault(row.main_sku, []).append(row)
    return {main: tuple(items) for main, items in grouped.items()}


def _merge_document(
    current: ProductFamilyDocument | None, incoming: tuple[_ImportRow, ...]
) -> ProductFamilyDocument:
    built = build_documents(row.row for row in incoming)[0]
    if current is None:
        return built
    incoming_children = {normalize_compare(child.sku): child for child in built.children}
    children = [
        incoming_children.pop(normalize_compare(child.sku), child)
        for child in current.children
    ]
    children.extend(incoming_children[key] for key in sorted(incoming_children))
    return replace(current, children=tuple(children))


def _write_document(connection: sqlite3.Connection, document: ProductFamilyDocument) -> None:
    connection.execute(
        """INSERT INTO documents(main_sku, document_json) VALUES (?, ?)
           ON CONFLICT(main_sku) DO UPDATE SET document_json = excluded.document_json""",
        (document.main_sku, _json(document.to_index_dict())),
    )
    connection.execute("DELETE FROM children WHERE main_sku = ?", (document.main_sku,))
    connection.executemany(
        """INSERT INTO children(main_sku, sku, display_name, sales_status_raw, status_flags_json)
           VALUES (?, ?, ?, ?, ?)""",
        (
            (
                document.main_sku,
                child.sku,
                child.display_name,
                child.sales_status_raw,
                _json(list(child.status_flags)),
            )
            for child in document.children
        ),
    )
    connection.execute("DELETE FROM product_fts WHERE main_sku = ?", (document.main_sku,))
    if document.searchable:
        categories = (
            *document.leaf_categories,
            *(part for path in document.category_paths for part in path),
        )
        connection.execute(
            "INSERT INTO product_fts(main_sku, cn_names, categories, en_aliases) VALUES (?, ?, ?, ?)",
            (
                document.main_sku,
                _fts_text(document.cn_names),
                _fts_text(categories),
                _fts_text(document.en_aliases),
            ),
        )


def _review_evidence(
    main_sku: str,
    rows: tuple[_ImportRow, ...],
    old_source: dict[str, dict[str, Any]],
) -> dict[str, list[str]]:
    evidence: dict[str, set[str]] = {"cn_names": set(), "english": set(), "categories": set()}
    incoming_children = {row.child_sku for row in rows}
    for old_child, old in old_source.items():
        if (old["main_sku"] == main_sku and old_child not in incoming_children
                and not is_excluded_product(main_sku, old_child)):
            _add_evidence(evidence, json.loads(old["evidence_json"]))
    for row in rows:
        if not is_excluded_product(main_sku, row.child_sku):
            _add_evidence(evidence, row.raw)
    return {key: sorted(values) for key, values in evidence.items()}


def _add_evidence(target: dict[str, set[str]], values: dict[str, str]) -> None:
    for key, bucket in (
        ("商品名称", "cn_names"),
        ("英文名称", "english"),
        ("英文关键字", "english"),
        ("商品目录", "categories"),
        ("商品一级目录", "categories"),
        ("商品二级目录", "categories"),
        ("商品三级目录", "categories"),
        ("商品四级目录", "categories"),
    ):
        value = clean_optional_text(values.get(key, ""))
        if value is not None:
            target[bucket].add(value)


def _has_prior_source(
    main_sku: str, old_source: dict[str, dict[str, Any]]
) -> bool:
    return any(old["main_sku"] == main_sku for old in old_source.values())


def _needs_review(
    current: ProductFamilyDocument | None,
    merged: ProductFamilyDocument,
    rows: tuple[_ImportRow, ...],
    old_source: dict[str, dict[str, Any]],
    evidence_hash: str,
) -> bool:
    rows = tuple(row for row in rows if not is_excluded_product(row.main_sku, row.child_sku))
    if not rows:
        return False
    if is_excluded_main_sku(merged.main_sku):
        return False
    if not any(clean_optional_text(value) for row in rows
               for value in (row.row.product_name, row.row.english_name, row.row.english_keywords)):
        return False
    if current is None:
        return True
    if _has_prior_source(merged.main_sku, old_source):
        return any(
            old_source.get(row.child_sku, {}).get("evidence_hash") != row.evidence_hash
            for row in rows
        )
    row_names = {clean_optional_text(row.row.product_name) for row in rows}
    row_names.discard(None)
    current_child_names = {clean_optional_text(child.display_name) for child in current.children}
    current_child_names.discard(None)
    new_names = not row_names.issubset(current_child_names)
    built = build_documents(row.row for row in rows)[0]
    new_aliases = set(built.en_aliases) - set(current.en_aliases)
    new_categories = set(built.leaf_categories) - set(current.leaf_categories)
    return bool(new_names or new_aliases or new_categories)


def _refresh_reviews(
    connection: sqlite3.Connection,
    current: dict[str, ProductFamilyDocument],
    merged: dict[str, ProductFamilyDocument],
    rows: tuple[_ImportRow, ...],
    old_source: dict[str, dict[str, Any]],
) -> int:
    pending = {row[0] for row in connection.execute("SELECT main_sku FROM daily_name_review WHERE status = 'pending'")}
    for main_sku, incoming in _group_by_main(rows).items():
        evidence = _review_evidence(main_sku, incoming, old_source)
        evidence_hash = _hash_json(evidence)
        if _needs_review(current.get(main_sku), merged[main_sku], incoming, old_source, evidence_hash):
            row = connection.execute(
                "SELECT evidence_hash, status FROM daily_name_review WHERE main_sku = ?",
                (main_sku,),
            ).fetchone()
            if row is None or row != (evidence_hash, "reviewed"):
                pending.add(main_sku)
    return len(pending)


def _write_review_rows(
    connection: sqlite3.Connection,
    current: dict[str, ProductFamilyDocument],
    merged: dict[str, ProductFamilyDocument],
    rows: tuple[_ImportRow, ...],
    old_source: dict[str, dict[str, Any]],
) -> None:
    for main_sku, incoming in _group_by_main(rows).items():
        evidence = _review_evidence(main_sku, incoming, old_source)
        evidence_hash = _hash_json(evidence)
        if not _needs_review(current.get(main_sku), merged[main_sku], incoming, old_source, evidence_hash):
            continue
        old = connection.execute(
            "SELECT evidence_hash, status FROM daily_name_review WHERE main_sku = ?",
            (main_sku,),
        ).fetchone()
        if old == (evidence_hash, "reviewed"):
            continue
        connection.execute(
            """INSERT INTO daily_name_review(main_sku, evidence_hash, evidence_json, status)
               VALUES (?, ?, ?, 'pending')
               ON CONFLICT(main_sku) DO UPDATE SET
                   evidence_hash = excluded.evidence_hash,
                   evidence_json = excluded.evidence_json,
                   status = 'pending'""",
            (main_sku, evidence_hash, _json(evidence)),
        )
