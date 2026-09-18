from __future__ import annotations

from collections.abc import Sequence
from datetime import UTC, datetime
from hashlib import sha256
import json
import os
from pathlib import Path
import subprocess
import sys

import numpy as np
import pytest
from openpyxl import Workbook

from store_scenario_inspiration.catalog.quality import QualityGateError
from store_scenario_inspiration.catalog.storage import CatalogStore
from store_scenario_inspiration.catalog.versioning import (
    CLEANING_RULES_VERSION,
    DOCUMENT_SCHEMA_VERSION,
    CatalogIndexError,
    CatalogIndexManager,
)
import store_scenario_inspiration.catalog.versioning as versioning


HEADERS = (
    "sku",
    "主SKU",
    "商品名称",
    "英文名称",
    "英文关键字",
    "销售状态",
    "商品目录",
    "商品一级目录",
    "商品二级目录",
    "商品三级目录",
    "商品四级目录",
)


def _write_source(path: Path, rows: tuple[tuple[str, ...], ...]) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "产品列表"
    worksheet.append(HEADERS)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    return path


def _write_two_sheet_source(path: Path) -> Path:
    workbook = Workbook()
    first = workbook.active
    first.title = "Sheet A"
    second = workbook.create_sheet("Sheet B")
    for sheet, row in (
        (first, _row("A-CHILD", "A-MAIN", "A款露营灯")),
        (second, _row("B-CHILD", "B-MAIN", "B款太阳能灯")),
    ):
        sheet.append(HEADERS)
        sheet.append(row)
    workbook.save(path)
    return path


def _row(sku: str, main_sku: str, name: str) -> tuple[str, ...]:
    return (
        sku,
        main_sku,
        name,
        "",
        "",
        "",
        "户外灯",
        "运动",
        "露营",
        "照明",
        "户外灯",
    )


class RecordingProvider:
    dimension = 2

    def __init__(self, model_id: str = "embedding-test-v1", *, fail: bool = False) -> None:
        self.model_id = model_id
        self.fail = fail
        self.calls: list[tuple[str, ...]] = []

    def embed(self, texts: Sequence[str]) -> np.ndarray:
        self.calls.append(tuple(texts))
        if self.fail:
            raise RuntimeError("provider unavailable")
        return np.asarray([[len(text), 1] for text in texts], dtype=np.float32)


def _version_dirs(root: Path) -> tuple[Path, ...]:
    versions = root / "versions"
    if not versions.exists():
        return ()
    return tuple(sorted(path for path in versions.iterdir() if not path.name.startswith(".staging-")))


def _snapshot(root: Path) -> dict[str, bytes]:
    return {
        str(path.relative_to(root)): path.read_bytes()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def test_rebuild_records_missing_identity_without_inventing_a_main_sku(tmp_path: Path) -> None:
    source = _write_source(tmp_path / "source.xlsx", (
        _row("CHILD-1", "MAIN-1", "露营灯"),
        _row("CHILD-2", "", "没有主SKU的灯"),
    ))
    manager = CatalogIndexManager(tmp_path / "catalog")
    manifest = manager.rebuild(source, sheet_name=None, provider=None)
    folder = manager.versions_dir / manifest.version_id
    imported = json.loads((folder / "source-import.json").read_text(encoding="utf-8"))
    assert imported["source_row_count"] == 2
    assert imported["indexed_source_row_count"] == 1
    assert imported["unassigned_rows"][0]["sku"] == "CHILD-2"
    assert imported["unassigned_rows"][0]["main_sku"] == ""
    assert manifest.document_count == 1
    assert manager.active_manifest().version_id == manifest.version_id


def test_failed_rebuild_keeps_active_version(tmp_path: Path) -> None:
    valid_source = _write_source(
        tmp_path / "valid.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    invalid_source = _write_source(
        tmp_path / "invalid.xlsx",
        (
            _row("DUPLICATE", "MAIN-1", "露营灯"),
            _row("DUPLICATE", "MAIN-2", "太阳能灯"),
        ),
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(valid_source, sheet_name=None, provider=None)

    with pytest.raises(QualityGateError):
        manager.rebuild(invalid_source, sheet_name=None, provider=None)

    active = manager.active_manifest()
    assert active is not None
    assert active.version_id == first.version_id


def test_rollback_swaps_active_and_previous(tmp_path: Path) -> None:
    source_v1 = _write_source(
        tmp_path / "v1.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    source_v2 = _write_source(
        tmp_path / "v2.xlsx", (_row("SKU-2", "MAIN-2", "太阳能灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(source_v1, sheet_name=None, provider=None)
    second = manager.rebuild(source_v2, sheet_name=None, provider=None)
    assert json.loads((manager.index_root / "active.json").read_text(encoding="utf-8")) == {
        "version_id": second.version_id,
        "previous_version_id": first.version_id,
    }

    restored = manager.rollback()

    assert restored.version_id == first.version_id
    assert json.loads((manager.index_root / "active.json").read_text(encoding="utf-8")) == {
        "version_id": first.version_id,
        "previous_version_id": second.version_id,
    }
    assert manager.rollback().version_id == second.version_id


def test_unchanged_build_skips_without_reading_workbook_or_changing_index(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(source, sheet_name=None, provider=None)
    before = _snapshot(manager.index_root)

    def fail_if_read(*args: object, **kwargs: object) -> object:
        raise AssertionError("unchanged rebuild must not read the workbook")

    monkeypatch.setattr(versioning, "iter_source_rows", fail_if_read)
    skipped = manager.rebuild(source, sheet_name=None, provider=None)

    assert skipped == first
    assert manager.last_rebuild_skipped is True
    assert _snapshot(manager.index_root) == before
    version = manager.index_root / "versions" / first.version_id
    assert not (version / "vectors.npy").exists()
    assert not (version / "vector-rows.json").exists()
    assert json.loads((version / "quality.json").read_text(encoding="utf-8"))["errors"] == []


def test_sheet_selector_is_part_of_build_identity_and_manifest(tmp_path: Path) -> None:
    source = _write_two_sheet_source(tmp_path / "two-sheets.xlsx")
    manager = CatalogIndexManager(tmp_path / "catalog")

    sheet_a = manager.rebuild(source, sheet_name="Sheet A", provider=None)
    sheet_b = manager.rebuild(source, sheet_name="Sheet B", provider=None)
    repeated_b = manager.rebuild(source, sheet_name="Sheet B", provider=None)

    assert sheet_a.source_sheet_name == "Sheet A"
    assert sheet_b.source_sheet_name == "Sheet B"
    assert sheet_b.source_modified_at.endswith("Z")
    assert sheet_b.version_id != sheet_a.version_id
    assert manager.last_rebuild_skipped is True
    assert repeated_b == sheet_b
    with CatalogStore.open_readonly(manager.active_store_path()) as store:
        assert store.get_document("B-MAIN") is not None
        assert store.get_document("A-MAIN") is None


def test_auto_sheet_selector_is_recorded_as_none(tmp_path: Path) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )

    manifest = CatalogIndexManager(tmp_path / "catalog").rebuild(
        source, sheet_name=None, provider=None
    )

    assert manifest.source_sheet_name is None
    assert manifest.source_modified_at.endswith("Z")


def test_schema_cleaning_rule_and_embedding_model_identity_changes_force_rebuild(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(source, sheet_name=None, provider=None)

    monkeypatch.setattr(versioning, "DOCUMENT_SCHEMA_VERSION", "3")
    schema_changed = manager.rebuild(source, sheet_name=None, provider=None)
    next_rules_version = str(int(versioning.CLEANING_RULES_VERSION) + 1)
    monkeypatch.setattr(versioning, "CLEANING_RULES_VERSION", next_rules_version)
    rules_changed = manager.rebuild(source, sheet_name=None, provider=None)
    provider = RecordingProvider()
    vector_build = manager.rebuild(source, sheet_name=None, provider=provider)

    assert schema_changed.version_id != first.version_id
    assert schema_changed.schema_version == "3"
    assert rules_changed.version_id != schema_changed.version_id
    assert rules_changed.cleaning_rules_version == next_rules_version
    assert vector_build.version_id != rules_changed.version_id
    assert vector_build.embedding_model_id == provider.model_id
    assert vector_build.vector_status == "present"
    vector_dir = manager.index_root / "versions" / vector_build.version_id
    assert (vector_dir / "vectors.npy").is_file()
    assert json.loads((vector_dir / "vector-rows.json").read_text(encoding="utf-8")) == [
        "MAIN-1"
    ]
    assert provider.calls == [("商品名称：露营灯",)]


@pytest.mark.parametrize("model_id", ["", "   ", "none", "NONE", " none "])
def test_provider_model_id_must_be_nonempty_and_cannot_impersonate_none(
    tmp_path: Path, model_id: str
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")

    with pytest.raises(ValueError, match="model_id"):
        manager.rebuild(source, sheet_name=None, provider=RecordingProvider(model_id))

    assert not manager.index_root.exists()


def test_provider_failure_cleans_staging_and_keeps_active_version(tmp_path: Path) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(source, sheet_name=None, provider=None)

    with pytest.raises(RuntimeError, match="provider unavailable"):
        manager.rebuild(source, sheet_name=None, provider=RecordingProvider(fail=True))

    assert manager.active_manifest() == first
    assert list((manager.index_root / "versions").glob(".staging-*")) == []


def test_source_change_during_build_is_rejected_before_activation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    hashes = iter(("a" * 64, "b" * 64))
    monkeypatch.setattr(versioning, "_sha256_file", lambda path: next(hashes))

    with pytest.raises(CatalogIndexError, match="changed during"):
        manager.rebuild(source, sheet_name=None, provider=None)

    assert manager.active_manifest() is None
    assert list((manager.index_root / "versions").glob(".staging-*")) == []
    assert _version_dirs(manager.index_root) == ()


def test_existing_same_second_version_id_advances_without_sleep(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    root = tmp_path / "catalog"
    fixed = datetime(2026, 8, 31, 10, 0, 0, tzinfo=UTC)
    monkeypatch.setattr(versioning, "_utc_now", lambda: fixed)
    source_hash = sha256(source.read_bytes()).hexdigest()
    identity = versioning._build_identity(
        source_hash,
        DOCUMENT_SCHEMA_VERSION,
        CLEANING_RULES_VERSION,
        "none",
        None,
        versioning._source_modified_at(source),
    )
    collision = root / "versions" / f"20260831T100000Z-{identity[:12]}"
    collision.mkdir(parents=True)

    manifest = CatalogIndexManager(root).rebuild(source, sheet_name=None, provider=None)

    assert manifest.version_id == f"20260831T100001Z-{identity[:12]}"


@pytest.mark.parametrize(
    "payload",
    ["{", '{"version_id":"../escape"}', '{"version_id":"safe","extra":1}'],
)
def test_active_manifest_rejects_corrupt_or_unsafe_pointer(
    tmp_path: Path, payload: str
) -> None:
    root = tmp_path / "catalog"
    root.mkdir()
    (root / "active.json").write_text(payload, encoding="utf-8")

    with pytest.raises(CatalogIndexError, match="pointer|version ID"):
        CatalogIndexManager(root).active_manifest()


def test_active_manifest_rejects_missing_or_corrupt_manifest(tmp_path: Path) -> None:
    root = tmp_path / "catalog"
    version_id = "20260831T100000Z-aaaaaaaaaaaa"
    (root / "versions" / version_id).mkdir(parents=True)
    (root / "active.json").write_text(
        json.dumps({"version_id": version_id, "previous_version_id": None}),
        encoding="utf-8",
    )
    manager = CatalogIndexManager(root)

    with pytest.raises(CatalogIndexError, match="manifest does not exist"):
        manager.active_manifest()

    (root / "versions" / version_id / "manifest.json").write_text("{}", encoding="utf-8")
    with pytest.raises(CatalogIndexError, match="manifest has invalid fields"):
        manager.active_manifest()


def test_rollback_without_previous_and_identical_pointers_are_explicit_errors(
    tmp_path: Path
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(source, sheet_name=None, provider=None)

    with pytest.raises(CatalogIndexError, match="previous pointer does not exist"):
        manager.rollback()

    (manager.index_root / "active.json").write_text(
        json.dumps(
            {
                "version_id": first.version_id,
                "previous_version_id": first.version_id,
            }
        ),
        encoding="utf-8",
    )
    with pytest.raises(CatalogIndexError, match="identical"):
        manager.rollback()


def test_active_state_failure_preserves_authoritative_pair_and_rollback_chain(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source_v1 = _write_source(
        tmp_path / "v1.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    source_v2 = _write_source(
        tmp_path / "v2.xlsx", (_row("SKU-2", "MAIN-2", "太阳能灯"),)
    )
    source_v3 = _write_source(
        tmp_path / "v3.xlsx", (_row("SKU-3", "MAIN-3", "充电露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(source_v1, sheet_name=None, provider=None)
    second = manager.rebuild(source_v2, sheet_name=None, provider=None)
    original_write_state = manager._write_active_state

    def fail_active(state: object) -> None:
        raise OSError("state unavailable")

    monkeypatch.setattr(manager, "_write_active_state", fail_active)
    with pytest.raises(OSError, match="state unavailable"):
        manager.rebuild(source_v3, sheet_name=None, provider=None)

    state = json.loads((manager.index_root / "active.json").read_text(encoding="utf-8"))
    assert state == {
        "version_id": second.version_id,
        "previous_version_id": first.version_id,
    }
    monkeypatch.setattr(manager, "_write_active_state", original_write_state)
    assert manager.rollback().version_id == first.version_id
    assert len(_version_dirs(manager.index_root)) == 3
    assert list((manager.index_root / "versions").glob(".staging-*")) == []


def test_derived_previous_failure_preserves_pair_and_rollback_chain(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sources = tuple(
        _write_source(
            tmp_path / f"v{number}.xlsx",
            (_row(f"SKU-{number}", f"MAIN-{number}", f"露营灯{number}"),),
        )
        for number in range(1, 4)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(sources[0], sheet_name=None, provider=None)
    second = manager.rebuild(sources[1], sheet_name=None, provider=None)
    original_write_previous = manager._write_derived_previous

    def fail_previous(version_id: str | None) -> None:
        raise OSError("derived previous unavailable")

    monkeypatch.setattr(manager, "_write_derived_previous", fail_previous)
    with pytest.raises(OSError, match="derived previous unavailable"):
        manager.rebuild(sources[2], sheet_name=None, provider=None)

    assert manager.active_manifest() == second
    state = json.loads((manager.index_root / "active.json").read_text(encoding="utf-8"))
    assert state["previous_version_id"] == first.version_id
    monkeypatch.setattr(manager, "_write_derived_previous", original_write_previous)
    assert manager.rollback().version_id == first.version_id


def test_rollback_state_failure_preserves_pair_and_can_be_retried(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source_v1 = _write_source(
        tmp_path / "v1.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    source_v2 = _write_source(
        tmp_path / "v2.xlsx", (_row("SKU-2", "MAIN-2", "太阳能灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(source_v1, sheet_name=None, provider=None)
    second = manager.rebuild(source_v2, sheet_name=None, provider=None)
    original_state = (manager.index_root / "active.json").read_bytes()
    original_write_state = manager._write_active_state
    monkeypatch.setattr(
        manager,
        "_write_active_state",
        lambda state: (_ for _ in ()).throw(OSError("rollback state unavailable")),
    )

    with pytest.raises(OSError, match="rollback state unavailable"):
        manager.rollback()

    assert (manager.index_root / "active.json").read_bytes() == original_state
    assert manager.active_manifest() == second
    monkeypatch.setattr(manager, "_write_active_state", original_write_state)
    assert manager.rollback().version_id == first.version_id


def test_schema_and_storage_failures_clean_staging_and_keep_active(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    valid = _write_source(
        tmp_path / "valid.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    invalid = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    workbook.active.append(("sku", "商品名称"))
    workbook.save(invalid)
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(valid, sheet_name=None, provider=None)

    with pytest.raises(ValueError, match="required headers"):
        manager.rebuild(invalid, sheet_name=None, provider=None)
    assert manager.active_manifest() == first
    assert list((manager.index_root / "versions").glob(".staging-*")) == []

    changed = _write_source(
        tmp_path / "changed.xlsx", (_row("SKU-2", "MAIN-2", "太阳能灯"),)
    )

    def fail_storage(*args: object, **kwargs: object) -> object:
        raise OSError("storage unavailable")

    monkeypatch.setattr(versioning.CatalogStore, "create", fail_storage)
    with pytest.raises(OSError, match="storage unavailable"):
        manager.rebuild(changed, sheet_name=None, provider=None)
    assert manager.active_manifest() == first
    assert list((manager.index_root / "versions").glob(".staging-*")) == []


def test_active_manifest_rejects_vector_status_artifact_mismatches(tmp_path: Path) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    absent = manager.rebuild(source, sheet_name=None, provider=None)
    absent_dir = manager.index_root / "versions" / absent.version_id
    (absent_dir / "vectors.npy").write_bytes(b"unexpected")

    with pytest.raises(CatalogIndexError, match="vector artifacts"):
        manager.active_manifest()

    (absent_dir / "vectors.npy").unlink()
    present = manager.rebuild(source, sheet_name=None, provider=RecordingProvider())
    present_dir = manager.index_root / "versions" / present.version_id
    (present_dir / "vector-rows.json").unlink()

    with pytest.raises(CatalogIndexError, match="vector artifacts"):
        manager.active_manifest()


def test_embedding_cache_is_reused_but_each_version_keeps_complete_vector_artifacts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    provider = RecordingProvider()
    first = manager.rebuild(source, sheet_name=None, provider=provider)
    first_dir = manager.index_root / "versions" / first.version_id

    monkeypatch.setattr(versioning, "CLEANING_RULES_VERSION", str(int(versioning.CLEANING_RULES_VERSION) + 1))
    second = manager.rebuild(source, sheet_name=None, provider=provider)
    second_dir = manager.index_root / "versions" / second.version_id

    assert provider.calls == [("商品名称：露营灯",)]
    assert (first_dir / "vectors.npy").read_bytes() == (second_dir / "vectors.npy").read_bytes()
    assert (first_dir / "vector-rows.json").read_bytes() == (
        second_dir / "vector-rows.json"
    ).read_bytes()
    assert list((manager.index_root / "embedding-cache" / "embeddings").glob("*.npy"))


def test_source_change_and_report_write_failure_keep_existing_active_and_clean_staging(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    first_source = _write_source(
        tmp_path / "first.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    changed_source = _write_source(
        tmp_path / "changed.xlsx", (_row("SKU-2", "MAIN-2", "太阳能灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(first_source, sheet_name=None, provider=None)
    actual_hash = sha256(changed_source.read_bytes()).hexdigest()
    hashes = iter((actual_hash, "f" * 64))
    original_sha256_file = versioning._sha256_file

    def source_hash_only(path: Path) -> str:
        if Path(path) == changed_source:
            return next(hashes)
        return original_sha256_file(path)

    monkeypatch.setattr(versioning, "_sha256_file", source_hash_only)

    with pytest.raises(CatalogIndexError, match="changed during"):
        manager.rebuild(changed_source, sheet_name=None, provider=None)

    assert manager.active_manifest() == first
    assert list((manager.index_root / "versions").glob(".staging-*")) == []

    monkeypatch.setattr(versioning, "_sha256_file", lambda path: sha256(path.read_bytes()).hexdigest())
    original_write_json = versioning._write_json_fsync

    def fail_quality(path: Path, payload: object) -> None:
        if path.name == "quality.json":
            raise OSError("quality write failed")
        original_write_json(path, payload)

    monkeypatch.setattr(versioning, "_write_json_fsync", fail_quality)
    with pytest.raises(OSError, match="quality write failed"):
        manager.rebuild(changed_source, sheet_name=None, provider=None)
    assert manager.active_manifest() == first
    assert list((manager.index_root / "versions").glob(".staging-*")) == []


def test_build_identity_is_canonical_and_final_version_never_contains_source_workbook(
    tmp_path: Path
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    manifest = manager.rebuild(source, sheet_name=None, provider=None)
    encoded = json.dumps(
        {
            "cleaning_rules_version": CLEANING_RULES_VERSION,
            "embedding_model_id": "none",
            "schema_version": DOCUMENT_SCHEMA_VERSION,
            "source_sha256": sha256(source.read_bytes()).hexdigest(),
            "source_modified_at": manifest.source_modified_at,
            "source_sheet_name": None,
        },
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    expected_identity = sha256(encoded).hexdigest()
    version_dir = manager.index_root / "versions" / manifest.version_id

    assert manifest.version_id.endswith(f"-{expected_identity[:12]}")
    assert {path.name for path in version_dir.iterdir()} == {
        "artifacts.json",
        "catalog.sqlite3",
        "delta.json",
        "manifest.json",
        "quality.json",
        "source-import.json",
    }
    assert source.read_bytes().startswith(b"PK")


def test_second_manager_writer_is_rejected_until_process_lock_is_released(
    tmp_path: Path,
) -> None:
    source_v1 = _write_source(
        tmp_path / "v1.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    source_v2 = _write_source(
        tmp_path / "v2.xlsx", (_row("SKU-2", "MAIN-2", "太阳能灯"),)
    )
    root = tmp_path / "catalog"
    first_manager = CatalogIndexManager(root)
    second_manager = CatalogIndexManager(root)
    first_manager.rebuild(source_v1, sheet_name=None, provider=None)

    with first_manager._writer_lock():
        with pytest.raises(CatalogIndexError, match="writer lock"):
            second_manager.rebuild(source_v2, sheet_name=None, provider=None)

    second = second_manager.rebuild(source_v2, sheet_name=None, provider=None)
    assert second_manager.active_manifest() == second


def test_active_validation_rejects_tampered_complete_artifacts(tmp_path: Path) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    manifest = manager.rebuild(source, sheet_name=None, provider=None)
    version = manager.index_root / "versions" / manifest.version_id

    artifacts = json.loads((version / "artifacts.json").read_text(encoding="utf-8"))
    assert set(artifacts) == {
        "catalog.sqlite3",
        "manifest.json",
        "quality.json",
        "delta.json",
        "source-import.json",
    }
    (version / "quality.json").write_text("{}", encoding="utf-8")

    with pytest.raises(CatalogIndexError, match="artifact hash mismatch"):
        manager.active_manifest()


def test_active_validation_rejects_sqlite_manifest_and_count_tampering(tmp_path: Path) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    manifest = manager.rebuild(source, sheet_name=None, provider=None)
    version = manager.index_root / "versions" / manifest.version_id
    import sqlite3

    with sqlite3.connect(version / "catalog.sqlite3") as connection:
        connection.execute("DELETE FROM children")
    # Re-sign only the modified SQLite: relational validation must still reject it.
    artifacts_path = version / "artifacts.json"
    artifacts = json.loads(artifacts_path.read_text(encoding="utf-8"))
    artifacts["catalog.sqlite3"] = sha256((version / "catalog.sqlite3").read_bytes()).hexdigest()
    artifacts_path.write_text(json.dumps(artifacts), encoding="utf-8")

    with pytest.raises(CatalogIndexError, match="child count"):
        manager.active_manifest()


def test_active_validation_rejects_resigned_invalid_delta_schema(tmp_path: Path) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    manifest = manager.rebuild(source, sheet_name=None, provider=None)
    version = manager.index_root / "versions" / manifest.version_id
    (version / "delta.json").write_text('{"metrics":{}}', encoding="utf-8")
    artifacts_path = version / "artifacts.json"
    artifacts = json.loads(artifacts_path.read_text(encoding="utf-8"))
    artifacts["delta.json"] = sha256((version / "delta.json").read_bytes()).hexdigest()
    artifacts_path.write_text(json.dumps(artifacts), encoding="utf-8")

    with pytest.raises(CatalogIndexError, match="delta report"):
        manager.active_manifest()


def _fail_writer_unlock(monkeypatch: pytest.MonkeyPatch) -> None:
    if os.name == "nt":
        import msvcrt

        original = msvcrt.locking

        def fail_unlock(fd: int, mode: int, size: int) -> object:
            if mode == msvcrt.LK_UNLCK:
                raise OSError("unlock unavailable")
            return original(fd, mode, size)

        monkeypatch.setattr(msvcrt, "locking", fail_unlock)
    else:
        import fcntl

        original = fcntl.flock

        def fail_unlock(fd: int, mode: int) -> object:
            if mode == fcntl.LOCK_UN:
                raise OSError("unlock unavailable")
            return original(fd, mode)

        monkeypatch.setattr(fcntl, "flock", fail_unlock)


def test_rebuild_returns_committed_manifest_when_writer_unlock_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    _fail_writer_unlock(monkeypatch)

    manifest = manager.rebuild(source, sheet_name=None, provider=None)

    assert manager._read_active_state(required=True).version_id == manifest.version_id


def test_rollback_returns_restored_manifest_when_writer_unlock_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    first_source = _write_source(
        tmp_path / "first.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    second_source = _write_source(
        tmp_path / "second.xlsx", (_row("SKU-2", "MAIN-2", "太阳能灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    first = manager.rebuild(first_source, sheet_name=None, provider=None)
    manager.rebuild(second_source, sheet_name=None, provider=None)
    _fail_writer_unlock(monkeypatch)

    restored = manager.rollback()

    assert restored == first
    assert manager._read_active_state(required=True).version_id == first.version_id


def test_writer_unlock_failure_preserves_the_original_rebuild_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    invalid = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    workbook.active.append(("sku", "商品名称"))
    workbook.save(invalid)
    manager = CatalogIndexManager(tmp_path / "catalog")
    _fail_writer_unlock(monkeypatch)

    with pytest.raises(ValueError, match="required headers"):
        manager.rebuild(invalid, sheet_name=None, provider=None)


def test_real_subprocess_writer_lock_rejects_then_releases(tmp_path: Path) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    root = tmp_path / "catalog"
    manager = CatalogIndexManager(root)
    manager.rebuild(source, sheet_name=None, provider=None)
    program = (
        "from pathlib import Path; import sys; "
        "from store_scenario_inspiration.catalog.versioning import CatalogIndexManager; "
        "manager=CatalogIndexManager(Path(sys.argv[1])); "
        "lock=manager._writer_lock(); lock.__enter__(); print('ready', flush=True); "
        "sys.stdin.read(); lock.__exit__(None, None, None)"
    )
    child = subprocess.Popen(
        [sys.executable, "-c", program, str(root)],
        cwd=Path.cwd(),
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
    )
    try:
        assert child.stdout is not None
        assert child.stdout.readline().strip() == "ready"
        with pytest.raises(CatalogIndexError, match="writer lock"):
            manager.rebuild(source, sheet_name=None, provider=None)
        assert child.stdin is not None
        stdout, stderr = child.communicate("release", timeout=10)
        assert child.returncode == 0, stderr or stdout
    finally:
        if child.poll() is None:
            child.terminate()
            child.wait(timeout=10)
    assert manager.rebuild(source, sheet_name=None, provider=None).version_id


def test_bad_active_same_source_cannot_skip_and_resets_skip_flag(tmp_path: Path) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    manifest = manager.rebuild(source, sheet_name=None, provider=None)
    manager.rebuild(source, sheet_name=None, provider=None)
    assert manager.last_rebuild_skipped is True
    (manager.index_root / "versions" / manifest.version_id / "catalog.sqlite3").unlink()

    with pytest.raises(CatalogIndexError, match="incomplete"):
        manager.rebuild(source, sheet_name=None, provider=None)

    assert manager.last_rebuild_skipped is False


def test_resigned_corrupt_vector_rows_are_rejected_by_embedding_id_validation(
    tmp_path: Path,
) -> None:
    source = _write_source(
        tmp_path / "source.xlsx", (_row("SKU-1", "MAIN-1", "露营灯"),)
    )
    manager = CatalogIndexManager(tmp_path / "catalog")
    manifest = manager.rebuild(source, sheet_name=None, provider=RecordingProvider())
    version = manager.index_root / "versions" / manifest.version_id
    rows_path = version / "vector-rows.json"
    rows_path.write_text('["other"]', encoding="utf-8")
    artifacts_path = version / "artifacts.json"
    artifacts = json.loads(artifacts_path.read_text(encoding="utf-8"))
    artifacts["vector-rows.json"] = sha256(rows_path.read_bytes()).hexdigest()
    artifacts_path.write_text(json.dumps(artifacts), encoding="utf-8")

    with pytest.raises(CatalogIndexError, match="vector rows"):
        manager.active_manifest()
