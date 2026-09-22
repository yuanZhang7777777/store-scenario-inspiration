import json
import sqlite3

import pytest
from openpyxl import Workbook

from store_scenario_inspiration.catalog.pilot import (
    build_fts,
    child_main_map,
    filter_country,
    fuse_rankings,
    keyword_search,
    load_children,
    load_country_stock,
    load_inventory_context,
    read_country_stock,
)


def test_bilingual_fts_rrf_and_country_filter_keep_global_audit() -> None:
    products = {
        "A": {"main_sku": "A", "standard_name_cn": "露营折叠椅", "standard_name_en": "Camping Chair"},
        "B": {"main_sku": "B", "standard_name_cn": "户外折叠桌", "standard_name_en": "Outdoor Folding Table"},
    }
    connection = build_fts(products)
    try:
        assert keyword_search(connection, "cn", "露营椅", 5)[0][0] == "A"
        assert keyword_search(connection, "en", "folding table", 5)[0][0] == "B"
    finally:
        connection.close()

    global_candidates = fuse_rankings(
        products,
        [("cn_vector", "露营", [("A", 0.9), ("B", 0.8)]),
         ("en_keyword", "folding table", [("B", 1.0)])],
        10,
    )
    country_candidates = filter_country(global_candidates, {"B": 3.0}, 10)
    assert {row["main_sku"] for row in global_candidates} == {"A", "B"}
    assert [row["main_sku"] for row in country_candidates] == ["B"]
    assert next(row for row in global_candidates if row["main_sku"] == "A")["country_available"] is False

    weighted = fuse_rankings(
        products,
        [("cn_vector", "主词", [("A", 0.9)], 1.0),
         ("cn_vector", "扩写", [("B", 0.9)], 0.7)],
        10,
    )
    assert [row["main_sku"] for row in weighted] == ["A", "B"]


def test_country_stock_only_uses_asset_child_mapping(tmp_path) -> None:
    products = {
        "A": {"inventory_metadata_json": '{"child_skus":["A-1"]}'},
        "B": {"inventory_metadata_json": '{"child_skus":["B-1"]}'},
    }
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总表格"
    sheet.append(["子SKU", "主SKU", "国家", "库存中心库存", "公共池库存"])
    sheet.append(["A-1", "A", "菲律宾", 2, 3])
    sheet.append(["A-X", "A", "菲律宾", 100, 100])
    sheet.append(["B-1", "B", "泰国", 7, 8])
    path = tmp_path / "stock.xlsx"
    workbook.save(path)

    allowed = child_main_map(products)
    assert load_country_stock(path, "PH", allowed) == {"A": 5.0}
    assert load_inventory_context(path, "PH", allowed) == ("PH", "available", {"A": 5.0})


def test_the_stock_read_gives_the_children_beside_the_mains(tmp_path) -> None:
    """The sheet the manager reads lists the children under their货号, and both
    totals come out of the same pass over the 35MB workbook."""
    products = {"A": {"inventory_metadata_json": '{"child_skus":["A-1","A-2"]}'},
                "B": {"inventory_metadata_json": '{"child_skus":["B-1"]}'}}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总表格"
    sheet.append(["子SKU", "主SKU", "国家", "库存中心库存", "公共池库存"])
    sheet.append(["A-1", "A", "菲律宾", 2, 3])
    sheet.append(["A-2", "A", "菲律宾", 0, 0])
    sheet.append(["B-1", "B", "菲律宾", 4, 0])
    path = tmp_path / "stock.xlsx"
    workbook.save(path)

    by_main, by_child = read_country_stock(path, "PH", child_main_map(products))

    assert by_main == {"A": 5.0, "B": 4.0}
    # Only the children there is something to ship of: a child sitting at zero
    # is an absence, and the sheet prints that absence itself.
    assert by_child == {"A-1": 5.0, "B-1": 4.0}


def test_a_stock_row_naming_a_different_main_sku_is_skipped_and_reported(tmp_path) -> None:
    """The catalogue says A-1 belongs to A. A row that says otherwise is the two
    files disagreeing, and crediting the stock to either one would be a number
    nobody can trace — but dropping the whole read over it would cost the
    operator a run, and they cannot fix a file they cannot see the fault in."""
    products = {"A": {"inventory_metadata_json": '{"child_skus":["A-1","A-2"]}'}}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总表格"
    sheet.append(["子SKU", "主SKU", "国家", "库存中心库存", "公共池库存"])
    sheet.append(["A-1", "Z", "菲律宾", 2, 3])
    sheet.append(["A-2", "A", "菲律宾", 4, 0])
    path = tmp_path / "stock.xlsx"
    workbook.save(path)

    notes: list[str] = []
    by_main, _ = read_country_stock(path, "PH", child_main_map(products), notes)

    # The honest row still counts; the contradictory one is gone, not guessed at.
    assert by_main == {"A": 4.0}
    assert notes[0].startswith("库存文件里有 1 行")
    assert "第 2 行 A-1" in notes[1]


def test_a_whole_file_of_bad_rows_is_a_count_rather_than_a_wall_of_text(tmp_path) -> None:
    """The note is read on the page. One line per broken row would bury the
    number that says how widespread the damage is."""
    children = [f"A-{index + 1}" for index in range(8)]
    products = {"A": {"inventory_metadata_json": json.dumps({"child_skus": children})}}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总表格"
    sheet.append(["子SKU", "主SKU", "国家", "库存中心库存", "公共池库存"])
    for child in children:
        sheet.append([child, "Z", "菲律宾", 1, 0])
    path = tmp_path / "stock.xlsx"
    workbook.save(path)

    notes: list[str] = []
    read_country_stock(path, "PH", child_main_map(products), notes)

    assert notes[0].startswith("库存文件里有 8 行")
    assert notes[-1] == "其余 3 行同样跳过。"


def test_a_products_children_are_read_from_the_catalogue_with_their_names(tmp_path) -> None:
    """The child list is what the product library holds, not what the stock
    workbook happened to contain: a child with no stock this month is still a
    child the operator may need to order."""
    path = tmp_path / "catalog.sqlite3"
    connection = sqlite3.connect(path)
    connection.execute("""CREATE TABLE products (
        main_sku TEXT, active INTEGER, indexable INTEGER,
        inventory_metadata_json TEXT, raw_evidence_json TEXT)""")
    connection.execute("INSERT INTO products VALUES (?, 1, 1, ?, ?)", (
        "A", json.dumps({"child_skus": ["A-1", "A-2"]}),
        json.dumps([{"child_sku": "A-1", "raw": {"商品名称": "甲款蓝"}},
                    {"child_sku": "A-2", "raw": {"商品名称": "甲款红"}}]),
    ))
    connection.execute("INSERT INTO products VALUES (?, 1, 1, ?, ?)",
                       ("B", "{}", "[]"))
    connection.commit()
    connection.close()

    children = load_children(path, ["A", "B", "A"])

    assert children == {"A": [("A-1", "甲款蓝"), ("A-2", "甲款红")], "B": []}


def test_uncovered_country_keeps_inventory_unknown_without_reading_stock(tmp_path) -> None:
    country, coverage, stock = load_inventory_context(tmp_path / "missing.xlsx", "BR", {})
    assert (country, coverage, stock) == ("BR", "unavailable", None)

    candidates = [{"main_sku": "A"}]
    assert filter_country(candidates, stock, 10) == []
    assert candidates == [{
        "main_sku": "A", "country_available": None, "country_available_quantity": None,
    }]
