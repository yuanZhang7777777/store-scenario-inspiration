from openpyxl import Workbook
import pytest

from store_scenario_inspiration.catalog.stock import read_stock_snapshot


def test_stock_is_country_specific_shared_and_a_replacement_snapshot(tmp_path):
    source = tmp_path / "stock.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总表格"
    sheet.append(["商品图片", "子SKU", "主SKU", "国家", "海外仓可发", "销售员"])
    for row in [
        ["", "a", "A", "菲律宾", 2, "甲"],
        ["", "a", "A", "菲律宾", 2, "乙"],
        ["", "b", "B", "泰国", 1, "甲"],
        ["", "c", "C", "菲律宾", 0, "甲"],
        ["", "x", "1A0000", "菲律宾", 10, "甲"],
        ["汇总", None, None, None, 15, None],
    ]:
        sheet.append(row)
    workbook.save(source)
    result = read_stock_snapshot(source, {"a": "A", "b": "B", "c": "C", "x": "1A0000"})
    assert result["country_main_skus"] == {"PH": ["A"], "TH": ["B"]}
    assert result["available_child_count"] == 2
    sheet.cell(2, 5, 0)
    sheet.cell(3, 5, 0)
    workbook.save(source)
    next_result = read_stock_snapshot(source, {"a": "A", "b": "B", "c": "C", "x": "1A0000"})
    assert next_result["country_main_skus"] == {"PH": [], "TH": ["B"]}
    assert next_result["source_sha256"] != result["source_sha256"]


def test_new_stock_child_uses_exact_known_main_without_guessing(tmp_path):
    source = tmp_path / "stock.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总表格"
    sheet.append(["子SKU", "主SKU", "国家", "海外仓可发"])
    sheet.append(["new-child", "A", "菲律宾", 2])
    workbook.save(source)
    result = read_stock_snapshot(source, {"old-child": "A"})
    assert result["country_main_skus"] == {"PH": ["A"]}
    assert result["main_only_matched_child_count"] == 1
    with pytest.raises(ValueError, match="disagrees"):
        read_stock_snapshot(source, {"new-child": "B", "old-child": "A"})
    with pytest.raises(ValueError, match="missing from ERP"):
        read_stock_snapshot(source, {"old-child": "B"})


def test_unmatched_product_is_reported_without_blocking_known_country_stock(tmp_path):
    source = tmp_path / "stock.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总表格"
    sheet.append(["子SKU", "主SKU", "国家", "海外仓可发", "商品名称", "库存中心库存", "公共池库存"])
    sheet.append(["a", "A", "菲律宾", 3, "已知商品", 2, 1])
    sheet.append(["missing", "NEW", "菲律宾", 2, "源表商品名", 2, None])
    sheet.append(["b", "B", "泰国", 4, "泰国商品", 4, 0])
    workbook.save(source)
    result = read_stock_snapshot(source, {"a": "A", "b": "B"}, allow_unmatched=True)
    assert result["country_main_skus"] == {"PH": ["A"], "TH": ["B"]}
    assert result["available_child_count"] == 2
    assert result["erp_mapping_complete"] is False
    assert result["unmatched_positive_main_skus"] == ["NEW"]
    assert result["unmatched_positive_rows"] == [{
        "source_row": 3, "country": "PH", "main_sku": "NEW", "child_sku": "missing",
        "raw": {"子SKU": "missing", "主SKU": "NEW", "国家": "菲律宾", "海外仓可发": 2,
                "商品名称": "源表商品名", "库存中心库存": 2, "公共池库存": None},
        "reason": "main_sku_missing_from_erp",
    }]
    with pytest.raises(ValueError, match="disagrees"):
        read_stock_snapshot(source, {"a": "WRONG", "b": "B"}, allow_unmatched=True)
