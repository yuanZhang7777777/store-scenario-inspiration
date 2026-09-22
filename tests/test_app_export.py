import io
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

# Importing the fixture itself is what registers it here; the fakes it installs
# are the same ones the rest of the API tests run against.
from test_app_api import client, run_job, upload  # noqa: E402, F401

from store_scenario_inspiration.app.export import (  # noqa: E402
    LABEL,
    MAX_ROW_POINTS,
    NAVY,
    SECTION,
    child_rows,
    deduped_rows,
    notes_for,
    picked_by_sku,
    scene_blocks,
    summaries,
    workbook,
)


def candidate(sku: str, *, rank: int, name: str | None = None,
              available: bool | None = True, quantity: int | None = 12) -> dict:
    return {"rank": rank, "main_sku": sku, "standard_name_cn": name or f"{sku}中文",
            "standard_name_en": f"{sku}EN", "rrf_score": 0.5, "channels": ["cn_keyword"],
            "matched_queries": [sku], "country_available": available,
            "country_available_quantity": quantity, "rerank": "related"}


def retrieval(*roles: dict) -> dict:
    return {"schema": "store-retrieval-v1", "country": "TH", "inventory": "available",
            "scenes": list(roles)}


def role(scene_name: str, product: str, *candidates: dict, dropped: list[dict] | None = None) -> dict:
    return {"scene_name": scene_name, "product_cn": product, "product_en": "EN",
            "queries": {"cn": [product], "en": []}, "candidates": list(candidates),
            "dropped": dropped or []}


def analysis(*scenes: tuple[str, list[str]]) -> dict:
    return {
        "model": "deepseek-flash",
        "manager_summary": {"executive_conclusion": "结论", "business_opportunity": "机会",
                            "recommended_actions": ["一", "二", "三"], "decision_boundary": "边界"},
        "store_profile": {"judgement": "店铺画像判断", "evidence": "画像依据"},
        "current_product_structure": {"judgement": "当前结构判断", "evidence": "当前依据"},
        "future_product_structure": {"judgement": "未来结构判断", "evidence": "未来依据",
                                     "priority_order": ["甲", "乙", "丙"]},
        "audiences": [{"audience_name": "租房上班族", "description": "25-35 岁", "evidence": "依据"}],
        "operation_strategy": [{"strategy_name": "打包卖", "description": "桌布配桌垫", "evidence": "依据"}],
        "scenes": [{"scene_name": name, "audience": "人群", "user_need": "需求", "evidence": "依据",
                    "product_needs": [{"product_cn": product, "product_en": "EN", "purpose": "用途"}
                                      for product in products]}
                   for name, products in scenes],
    }


def pick(scene_name: str, product: str, sku: str) -> dict:
    return {"scene_name": scene_name, "product_cn": product, "main_sku": sku}


def blocks(analysis_doc, retrieval_doc, picks) -> dict[str, list[list]]:
    """The scene board as a reader sees it: scene name → its rows."""
    return dict(scene_blocks(analysis_doc, retrieval_doc, picks))


# ---------- the scene board ----------

def test_a_scene_is_a_block_of_columns_and_its_skus_are_the_rows() -> None:
    """The manager copies a column into the ERP. A cell full of line breaks made
    that impossible, and it hid the stock order inside one cell."""
    board = blocks(
        analysis(("小户型客厅", ["茶几桌布"])),
        retrieval(role("小户型客厅", "茶几桌布",
                       candidate("XX119", rank=1), candidate("XX4", rank=2))),
        [pick("小户型客厅", "茶几桌布", "XX119"), pick("小户型客厅", "茶几桌布", "XX4")],
    )

    assert list(board) == ["小户型客厅"]
    assert board["小户型客厅"] == [["XX119", "XX119中文", 12.0], ["XX4", "XX4中文", 12.0]]


def test_the_best_stocked_sku_sits_at_the_top_of_its_scene() -> None:
    """At the top of a ranked list the candidates are about as relevant as each
    other; what separates them is whether there are 2 of them or 200."""
    board = blocks(
        analysis(("客厅", ["桌布"])),
        retrieval(role("客厅", "桌布",
                       candidate("A", rank=1, quantity=3), candidate("B", rank=9, quantity=9))),
        [pick("客厅", "桌布", "A"), pick("客厅", "桌布", "B")],
    )

    assert [row[0] for row in board["客厅"]] == ["B", "A"]


def test_when_the_stock_was_never_read_the_recall_order_stands() -> None:
    """Every quantity is None, so the four channels' own ranking is the only
    ordering there is any basis for."""
    board = blocks(
        analysis(("客厅", ["桌布"])),
        retrieval(role("客厅", "桌布",
                       candidate("A", rank=2, available=None, quantity=None),
                       candidate("B", rank=1, available=None, quantity=None))),
        [pick("客厅", "桌布", "A"), pick("客厅", "桌布", "B")],
    )

    assert [row[0] for row in board["客厅"]] == ["B", "A"]


def test_a_scene_nobody_picked_from_is_left_out_rather_than_drawn_empty() -> None:
    board = blocks(
        analysis(("客厅", ["桌布"]), ("卧室", ["床旗"])),
        retrieval(role("客厅", "桌布", candidate("A", rank=1)),
                  role("卧室", "床旗", candidate("B", rank=1))),
        [pick("客厅", "桌布", "A")],
    )

    assert list(board) == ["客厅"]


def test_one_sku_reached_by_two_roles_is_one_row_in_its_scene() -> None:
    """Two roles pointing at the same货号 is one product, and the sheet is a list
    to buy from — not a log of how many roles happened to name it."""
    board = blocks(
        analysis(("客厅", ["桌布", "茶几垫"])),
        retrieval(role("客厅", "桌布", candidate("A", rank=5)),
                  role("客厅", "茶几垫", candidate("A", rank=2))),
        [pick("客厅", "桌布", "A"), pick("客厅", "茶几垫", "A")],
    )

    assert board["客厅"] == [["A", "A中文", 12.0]]


def test_a_sku_that_was_never_recalled_is_dropped_rather_than_trusted() -> None:
    """The picks arrive from the browser. Names and stock are read off disk, so a
    row the run never produced cannot be written out at all."""
    board = blocks(
        analysis(("客厅", ["桌布"])),
        retrieval(role("客厅", "桌布", candidate("A", rank=1))),
        [pick("客厅", "桌布", "A"), pick("客厅", "桌布", "不存在的SKU")],
    )

    assert [row[0] for row in board["客厅"]] == ["A"]


def test_a_sku_the_verdict_pass_dropped_is_still_exported_when_it_was_picked() -> None:
    """A pick can be a row the operator fished back out of the dropped pile."""
    board = blocks(
        analysis(("客厅", ["驱蚊液"])),
        retrieval(role("客厅", "驱蚊液", candidate("A", rank=1),
                       dropped=[candidate("Z", rank=31, name="捕虫器")])),
        [pick("客厅", "驱蚊液", "Z")],
    )

    assert board["客厅"] == [["Z", "捕虫器", 12.0]]


# ---------- the stock column ----------

def test_out_of_stock_says_so_rather_than_showing_a_zero() -> None:
    """"0" reads as a number somebody could top up; "没货" reads as a fact about
    the country."""
    board = blocks(
        analysis(("客厅", ["桌布"])),
        retrieval(role("客厅", "桌布",
                       candidate("A", rank=1, available=False, quantity=0),
                       candidate("B", rank=2, available=True, quantity=40))),
        [pick("客厅", "桌布", "A"), pick("客厅", "桌布", "B")],
    )

    assert [row[2] for row in board["客厅"]] == [40.0, "没货"]


def test_a_count_is_written_as_a_number_so_the_column_can_be_sorted() -> None:
    """The stock reader hands whole numbers over as floats, so 1.0 used to fall
    through to a bare "有货" and the column lost the number it exists for. It now
    lands in the sheet as a number, which is what makes the column sortable and
    filterable in Excel instead of a column of text that only looks sorted."""
    board = blocks(
        analysis(("客厅", ["桌布"])),
        retrieval(role("客厅", "桌布", candidate("A", rank=1, quantity=1.0),
                       candidate("B", rank=2, quantity=6.0),
                       candidate("C", rank=3, quantity=2.5))),
        [pick("客厅", "桌布", sku) for sku in ("A", "B", "C")],
    )

    assert sorted(row[2] for row in board["客厅"]) == [1.0, 2.5, 6.0]


def test_a_store_whose_stock_was_never_read_is_not_called_out_of_stock() -> None:
    """No answer to "有没有货" is a different thing from "没有货", and printing
    the second when we mean the first is a lie about the country."""
    board = blocks(
        analysis(("客厅", ["桌布"])),
        retrieval(role("客厅", "桌布", candidate("A", rank=1, available=None, quantity=None))),
        [pick("客厅", "桌布", "A")],
    )

    assert board["客厅"][0][2] == "未读到库存表"


# ---------- the sub-SKU sheet ----------

def test_every_child_is_listed_with_its_own_quantity() -> None:
    """The question this sheet answers is which child the goods are actually in,
    so the children this country cannot ship are the rows that carry the answer."""
    rows = child_rows(
        picked_by_sku(analysis(("客厅", ["桌布"])),
                      retrieval(role("客厅", "桌布", candidate("A", rank=1))),
                      [pick("客厅", "桌布", "A")]),
        {"A": [("A-2", "甲二号"), ("A-1", "甲一号")]},
        {"A-1": 5.0},
    )

    assert rows == [["A", "A中文", "A-1", "甲一号", 5.0, "客厅 · 桌布"],
                    ["A", "A中文", "A-2", "甲二号", 0.0, "客厅 · 桌布"]]


def test_a_product_the_catalogue_holds_no_child_for_still_gets_a_line() -> None:
    """A gap where a货号 should be reads as an omission rather than as a fact
    about the product library."""
    rows = child_rows(
        picked_by_sku(analysis(("客厅", ["桌布"])),
                      retrieval(role("客厅", "桌布", candidate("A", rank=1))),
                      [pick("客厅", "桌布", "A")]),
        {"A": []},
        {"A-1": 5.0},
    )

    assert rows == [["A", "A中文", "—", "产品库中未登记子款", 0.0, "客厅 · 桌布"]]


def test_a_country_we_could_not_read_says_so_instead_of_zero() -> None:
    rows = child_rows(
        picked_by_sku(analysis(("客厅", ["桌布"])),
                      retrieval(role("客厅", "桌布", candidate("A", rank=1))),
                      [pick("客厅", "桌布", "A")]),
        {"A": [("A-1", "甲一号")]},
        None,
    )

    assert rows[0][4] == "未核验"


# ---------- the buy-list ----------

def test_a_sku_picked_in_several_places_is_one_line_that_says_where() -> None:
    """The scene board repeats a SKU once per scene, which is what it is for. A
    list to upload from wants the opposite, and one SKU carrying two scenes is
    worth seeing rather than dissolving."""
    rows = deduped_rows(picked_by_sku(
        analysis(("客厅", ["桌布"]), ("卧室", ["床旗"])),
        retrieval(role("客厅", "桌布", candidate("A", rank=1)),
                  role("卧室", "床旗", candidate("A", rank=1))),
        [pick("客厅", "桌布", "A"), pick("卧室", "床旗", "A")],
    ))

    assert rows == [["A", "A中文", 12.0, "客厅 · 桌布\n卧室 · 床旗"]]


def test_the_buy_list_only_names_the_places_that_were_actually_picked() -> None:
    rows = deduped_rows(picked_by_sku(
        analysis(("客厅", ["桌布", "茶几垫"])),
        retrieval(role("客厅", "桌布", candidate("A", rank=1)),
                  role("客厅", "茶几垫", candidate("A", rank=1))),
        [pick("客厅", "桌布", "A")],
    ))

    assert rows == [["A", "A中文", 12.0, "客厅 · 桌布"]]


def test_the_buy_list_counts_a_dropped_row_that_was_picked_the_same_way() -> None:
    rows = deduped_rows(picked_by_sku(
        analysis(("客厅", ["驱蚊液"])),
        retrieval(role("客厅", "驱蚊液", candidate("A", rank=1, name="驱蚊液"),
                       dropped=[candidate("Z", rank=31, name="捕虫器")])),
        [pick("客厅", "驱蚊液", "Z")],
    ))

    assert rows == [["Z", "捕虫器", 12.0, "客厅 · 驱蚊液"]]


# ---------- the store line ----------

def test_the_store_line_leads_with_the_product_structure() -> None:
    """What the store makes its money on is what the person opening this sheet
    came for; the profile is the sentence that puts the rest in context."""
    found = dict(summaries(analysis(("小户型客厅", ["茶几桌布"])),
                           {"store_name": "Shopee-15005TH"}, "TH"))

    assert list(found)[:3] == ["店铺名称", "国家", "当前产品结构"]
    assert found["店铺名称"] == "Shopee-15005TH"
    assert found["国家"] == "泰国"
    assert found["当前产品结构"] == "当前结构判断"
    assert found["未来产品结构"] == "未来结构判断\n优先顺序：甲、乙、丙"
    assert found["店铺画像"] == "店铺画像判断"
    assert found["人群"] == "租房上班族：25-35 岁"
    assert found["场景"] == "小户型客厅"
    assert found["未来运营策略"] == "打包卖：桌布配桌垫"


def test_the_conclusions_do_not_recite_the_evidence_back() -> None:
    """The sheet is read for what the report makes of the numbers. The support it
    was built on is already on the sheet the numbers live on, and repeating it
    here reads as a report that decided nothing."""
    found = dict(summaries(analysis(("客厅", ["桌布"])), {}, "TH"))

    for field in ("当前产品结构", "未来产品结构", "店铺画像"):
        assert "依据" not in found[field]


def test_the_country_is_named_the_way_the_manager_names_it() -> None:
    """The code is what the catalogue is keyed by and tells the reader nothing
    about which warehouse the goods sit in."""
    assert dict(summaries(analysis(), {}, "PH"))["国家"] == "菲律宾"


def test_a_report_with_no_audiences_or_strategies_leaves_those_cells_empty() -> None:
    bare = {**analysis(), "audiences": [], "operation_strategy": [], "scenes": []}

    found = dict(summaries(bare, {}, "PH"))

    assert found["人群"] == ""
    assert found["未来运营策略"] == ""
    assert found["场景"] == ""


def test_the_store_line_reads_across_rather_than_down() -> None:
    """Stacked one per row, eight conclusions are a screenful of tall rows and
    the table the sheet exists for is somewhere below them."""
    payload = workbook([], summaries(analysis(("客厅", ["桌布"])),
                                     {"store_name": "店"}, "TH"), [], country="TH",
                       children=[])
    sheet = load_workbook(io.BytesIO(payload))["店铺概况"]

    labels = [cell.value for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]]
    assert (labels[0], labels[-1]) == ("店铺名称", "未来运营策略")
    assert (values[0], values[1]) == ("店", "泰国")
    # One row of values, so nothing runs on below it.
    assert sheet["A3"].value is None


def test_a_long_conclusion_still_fits_inside_one_excel_row() -> None:
    """The columns are wide on purpose. A seven-hundred-character conclusion in
    a narrow one is a row Excel will not draw past 409 points, and the tail of
    it would simply be missing — which is what the operator saw."""
    payload = workbook([], [("店铺名称", "店"), ("国家", "泰国"), ("店铺画像", ""),
                            ("当前产品结构", ""), ("未来产品结构", "很长的结论" * 140),
                            ("人群", ""), ("场景", ""), ("未来运营策略", "")], [],
                       country="TH", children=[])
    sheet = load_workbook(io.BytesIO(payload))["店铺概况"]

    assert sheet.row_dimensions[2].height < MAX_ROW_POINTS


def test_no_row_is_ever_asked_for_more_height_than_excel_will_draw() -> None:
    """Excel refuses to draw a row past 409 points, so a very long product name
    would lose its tail behind the row's edge on a sheet that looked complete."""
    rows = [[f"SKU-{index:02d}", "很长很长的中文名称" * 3, "12"] for index in range(60)]

    payload = workbook([("客厅", rows)], [], [], country="TH", children=[])

    heights = [dim.height for dim in
               load_workbook(io.BytesIO(payload))["主SKU清单"].row_dimensions.values()
               if dim.height is not None]
    assert heights and max(heights) <= MAX_ROW_POINTS


# ---------- how the workbook is dressed to be read ----------

def two_scenes() -> list[tuple[str, list[list]]]:
    return [("客厅", [["A", "桌布", 12.0]]), ("卧室", [["B", "床旗", 3.0]])]


XMLNS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def styled_cells(payload: bytes) -> dict[int, list[bool]]:
    """Per row of the board, whether each cell the workbook wrote carries a style.

    Straight from the file: openpyxl's reader drops the style it wrote on the
    cells a merge covers, so reading it back through the cell objects would
    report a bordered header as bare.
    """
    xml = zipfile.ZipFile(io.BytesIO(payload)).read("xl/worksheets/sheet1.xml")
    rows = ET.fromstring(xml).find(f"{XMLNS}sheetData")
    return {int(row.get("r")): [cell.get("s") is not None for cell in row]
            for row in rows}


def test_two_scenes_are_two_tables_and_not_one_wide_one() -> None:
    """Flush against each other the last column of one scene and the first of the
    next read as one very wide table with a stray header in the middle."""
    sheet = load_workbook(io.BytesIO(workbook(two_scenes(), [], [], country="TH",
                                              children=[])))["主SKU清单"]

    # Scene one holds columns A-C and its blank, so scene two opens two columns
    # past where the eye expects it — that gap is what separates the blocks.
    assert sheet.cell(row=1, column=5).value == "卧室"
    assert sheet.cell(row=1, column=4).value is None
    # The gap is a real column, not a merged cell pretending to be one.
    assert sheet.column_dimensions["D"].width > 0


def test_the_scene_title_and_its_header_read_as_headers() -> None:
    sheet = load_workbook(io.BytesIO(workbook(two_scenes(), [], [], country="TH",
                                              children=[])))["主SKU清单"]

    title = sheet.cell(row=1, column=1)
    assert title.value == "客厅"
    assert title.fill.fgColor.rgb.endswith(NAVY)
    assert title.font.b is True
    header = sheet.cell(row=2, column=1)
    assert header.value == "主SKU"
    assert header.fill.fgColor.rgb.endswith(LABEL)
    assert header.font.b is True


def test_every_cell_of_the_board_is_bordered_so_a_row_can_be_followed() -> None:
    """The eye tracks across a very wide sheet; an unbordered row is the one that
    gets read one column off.

    Read from the file rather than through openpyxl, which drops the style it
    wrote on the cells a merge covers. The gutter columns between the blocks are
    left bare on purpose, so they are simply not written.
    """
    payload = workbook(two_scenes(), [], [], country="TH", children=[])

    for row, styled in styled_cells(payload).items():
        assert styled and all(styled), row


def test_a_count_lands_as_a_number_the_column_can_be_sorted_by() -> None:
    """Written as text the quantity column only looks sorted: Excel sorts "9"
    above "40". A number makes the column sort and filter like one."""
    sheet = load_workbook(io.BytesIO(workbook(two_scenes(), [], [], country="TH",
                                              children=[])))["主SKU清单"]

    cell = sheet.cell(row=3, column=3)
    assert cell.value == 12.0
    assert cell.number_format == "#,##0.###"


def test_the_flat_lists_keep_their_header_on_screen_and_on_the_page() -> None:
    """Sixty rows in, the header is the only thing that says which column is the
    stock and which is the child; scrolled or printed, it has to still be there."""
    rows = [[f"SKU-{index:02d}", f"名称{index}", "C-1", "子款", 3.0, "客厅 · 桌布"]
            for index in range(60)]
    sheet = load_workbook(io.BytesIO(workbook([], [], [], country="TH",
                                              children=rows)))["子SKU清单"]

    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F61"
    assert sheet.print_title_rows == "$1:$1"


def test_the_prose_sheets_hide_the_grid_so_it_stops_cutting_the_sentences() -> None:
    payload = workbook(two_scenes(), summaries(analysis(), {}, "TH"), [],
                       country="TH", children=[])
    book = load_workbook(io.BytesIO(payload))

    assert book["店铺概况"].sheet_view.showGridLines is False
    assert book["口径说明"].sheet_view.showGridLines is False
    # Left unset, which is how a sheet shows its grid — the lists are the sheets
    # the grid is for.
    assert book["子SKU清单"].sheet_view.showGridLines is not False


def test_a_group_heading_in_the_notes_is_a_bar_rather_than_one_more_line() -> None:
    """The notes run to thirty-odd lines in four groups. At one weight the group
    headings dissolve into the lines they are meant to separate."""
    sheet = load_workbook(io.BytesIO(workbook([], [], [("店铺", "店"), ("库存口径", ""),
                                                       ("库存快照文件", "x.xlsx")],
                                              country="TH", children=[])))["口径说明"]

    heading = sheet.cell(row=2, column=1)
    assert heading.value == "库存口径"
    assert heading.fill.fgColor.rgb.endswith(SECTION)
    assert heading.font.b is True


def test_every_page_says_which_store_and_which_day_it_came_from() -> None:
    """The sheet gets printed and forwarded; the file name does not travel with
    it, and a stack of these from three stores is otherwise indistinguishable."""
    sheet = load_workbook(io.BytesIO(workbook(
        [], [("店铺名称", "Shopee-15005TH"), ("国家", "泰国")], [], country="TH",
        children=[])))["主SKU清单"]

    assert "Shopee-15005TH" in sheet.oddHeader.left.text
    assert "泰国" in sheet.oddHeader.left.text
    assert "&P" in sheet.oddFooter.right.text


# ---------- how the sheet says it was made ----------

def snapshot(filename: str) -> dict:
    return {"filename": filename, "sha256": "a" * 64, "captured_at": "2026-09-20 10:00",
            "file_modified_at": "2026-09-19 09:30 (+08:00)"}


def test_the_notes_name_the_snapshot_this_result_was_built_from(tmp_path) -> None:
    stock = tmp_path / "库存明细.xlsx"
    stock.write_bytes(b"x")

    notes = dict(notes_for(
        retrieval={**retrieval(), "inventory_snapshot": snapshot("库存明细.xlsx")},
        params={"scene_count": 6}, store={"store_name": "店"},
        store_id="store-1", stock_path=stock, exported=3,
    ))

    assert notes["库存快照文件"] == "库存明细.xlsx"
    assert notes["库存快照时间"] == "2026-09-19 09:30 (+08:00)"
    assert notes["库存快照 SHA256"] == "a" * 64
    assert notes["导出条数"] == "3 个产品"
    assert notes["这次读到了吗"] == "读到了"
    assert notes["库存数量"].startswith("仅对应本次结果绑定的库存快照")


def test_a_file_that_appeared_later_never_stands_in_for_the_one_that_was_read(tmp_path) -> None:
    """The export is handed whatever stock file is configured *now*. That file is
    not the evidence for these rows, so borrowing its timestamp would date the
    result to a snapshot it was never built from."""
    stock = tmp_path / "后来放进去的.xlsx"
    stock.write_bytes(b"x")

    notes = dict(notes_for(
        retrieval={**retrieval(), "inventory": "unavailable"}, params={}, store={},
        store_id="store-1", stock_path=stock, exported=0,
    ))

    assert notes["库存快照时间"] == "历史结果未记录，请重新匹配后核验"
    assert notes["库存快照文件"] == "未记录"
    assert notes["这次读到了吗"] == "未核验，本次候选不提供库存承诺"


def test_the_snapshot_the_results_used_survives_the_file_being_gone(tmp_path) -> None:
    notes = dict(notes_for(
        retrieval={**retrieval(), "inventory_snapshot": snapshot("已删除的库存.xlsx")},
        params={}, store={}, store_id="store-1",
        stock_path=tmp_path / "没有这个文件.xlsx", exported=0,
    ))

    assert notes["库存快照文件"] == "已删除的库存.xlsx"
    assert notes["库存快照时间"] == "2026-09-19 09:30 (+08:00)"


# ---------- the download ----------

def sheet_of(response) -> dict:
    book = load_workbook(io.BytesIO(response.content))
    return {name: [[cell.value for cell in row] for row in book[name].iter_rows()]
            for name in book.sheetnames}


def test_the_export_arrives_as_a_named_workbook(client) -> None:
    store_id = upload(client).json()["id"]
    run_job(client, store_id)

    response = client.post(f"/api/stores/{store_id}/adoption/export",
                           json={"picks": [pick("庭院遮阳", "遮阳棚替换布", "SKU-1")]})

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    sheets = sheet_of(response)
    assert list(sheets) == ["主SKU清单", "子SKU清单", "店铺概况", "口径说明"]

    # The scene is the column, its SKUs are the rows, and the stock header
    # carries the country's name so the block reads on its own.
    board = sheets["主SKU清单"]
    assert board[0][0] == "庭院遮阳"
    assert board[1][:3] == ["主SKU", "名称", "菲律宾有货"]
    assert board[2][:3] == ["SKU-1", "遮阳网", "未读到库存表"]

    line = sheets["店铺概况"]
    assert line[0][0] == "店铺名称"
    assert line[1][0] == "Shopee-13021PH"
    assert line[0][2] == "当前产品结构"


def test_the_sub_sku_sheet_names_every_child_of_every_picked_sku(client) -> None:
    store_id = upload(client).json()["id"]
    run_job(client, store_id)

    response = client.post(f"/api/stores/{store_id}/adoption/export",
                           json={"picks": [pick("庭院遮阳", "遮阳棚替换布", "SKU-1")]})

    children = sheet_of(response)["子SKU清单"]
    assert children[0] == ["主SKU", "主SKU名称", "子SKU", "子SKU名称", "菲律宾有货数量", "用在哪个场景"]
    assert children[1] == ["SKU-1", "遮阳网", "SKU-1-1", "SKU-1子款", "未核验",
                           "庭院遮阳 · 遮阳棚替换布"]


def test_asking_for_dedupe_adds_the_buy_list_beside_the_board(client) -> None:
    """The operator chooses. Left alone the workbook is the board it always was;
    asked for, a second list keeps each SKU once and says where it was used."""
    store_id = upload(client).json()["id"]
    run_job(client, store_id)

    # The fake catalogue answers every role with the same SKU, which is exactly
    # the case worth collapsing: one product carrying two roles in one scene.
    response = client.post(f"/api/stores/{store_id}/adoption/export", json={
        "picks": [pick("庭院遮阳", "遮阳棚替换布", "SKU-1"),
                  pick("庭院遮阳", "风扇", "SKU-1")],
        "dedupe": True,
    })

    sheets = sheet_of(response)
    assert list(sheets) == ["主SKU清单", "子SKU清单", "店铺概况", "去重商品清单", "口径说明"]
    # One scene is one block of columns, so the two roles collapse to one row
    # there; the buy-list is where the two roles stay legible.
    assert [row[0] for row in sheets["主SKU清单"][2:] if row[0]] == ["SKU-1"]

    buy = sheets["去重商品清单"]
    assert buy[0] == ["主SKU", "主SKU产品名称", "菲律宾有货数量", "用在哪里"]
    assert buy[1] == ["SKU-1", "遮阳网", "未读到库存表",
                      "庭院遮阳 · 遮阳棚替换布\n庭院遮阳 · 风扇"]


def test_exporting_before_anything_was_picked_says_so_on_the_board(client) -> None:
    store_id = upload(client).json()["id"]
    run_job(client, store_id)

    response = client.post(f"/api/stores/{store_id}/adoption/export", json={"picks": []})

    assert response.status_code == 200
    assert sheet_of(response)["主SKU清单"][0][0] == "本次没有勾选任何商品"


def test_exporting_a_store_whose_recall_never_ran_says_which_stage_is_missing(client) -> None:
    store_id = upload(client).json()["id"]

    response = client.post(f"/api/stores/{store_id}/adoption/export", json={"picks": []})

    assert response.status_code == 409
    assert "retrieval.json" in response.json()["detail"]


def test_a_result_from_before_the_child_detail_existed_asks_for_a_rerun(client) -> None:
    """Rewriting the sheet changed what the run has to leave behind. A stored
    result without it cannot be exported, and the fix is a free local rerun —
    the message says which one."""
    store_id = upload(client).json()["id"]
    run_job(client, store_id)
    (client.app.state.workspace.path(store_id, "stock_children.json")).unlink()

    response = client.post(f"/api/stores/{store_id}/adoption/export", json={"picks": []})

    assert response.status_code == 409
    assert "子款库存明细" in response.json()["detail"]


def test_the_notes_sheet_names_the_parameters_this_run_used(client) -> None:
    store_id = upload(client).json()["id"]
    run_job(client, store_id, params={"scene_count": 6, "products_per_scene": 8,
                                      "expansion_terms": 2, "recall_limit": 5})

    response = client.post(f"/api/stores/{store_id}/adoption/export", json={"picks": []})
    notes = dict((row[0], row[1]) for row in sheet_of(response)["口径说明"])

    assert notes["写几个场景"] == "6"
    assert notes["每个商品找几个 SKU"] == "5"
    assert notes["店铺 ID"] == store_id
    assert notes["国家"] == "菲律宾"


def test_a_setting_chosen_from_a_list_reads_back_in_the_words_it_was_chosen_in() -> None:
    """"all" in the sheet tells the reader nothing about what the run did with
    out-of-stock candidates; the dropdown's own phrase does."""
    notes = dict(notes_for(
        retrieval={**retrieval(), "rerank": {"mode": "mark_only", "provider": "jev"}},
        params={"stock_filter": "all", "scene_count": 6},
        store={}, store_id="store-1", stock_path=Path("x"), exported=0,
    ))

    assert notes["没货的要不要留着"] == "都留着，标出有没有货"
    assert notes["实际用的模式"] == "不排除，都留着（只标出来）"
